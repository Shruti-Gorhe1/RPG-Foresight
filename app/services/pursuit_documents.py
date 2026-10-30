"""Extract cited, reviewable evidence from uploaded tender documents."""
from __future__ import annotations

import io
import re
from datetime import datetime
import zipfile
from typing import Any

from observability.tracing import traced_span


MAX_DOCUMENT_SIZE = 10 * 1024 * 1024
MAX_DOCUMENTS = 10
MAX_TOTAL_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS_PER_DOCUMENT = 200_000
MAX_PDF_PAGES = 500
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

_FINDING_TERMS = {
    "eligibility": (
        "eligibility", "eligible", "ineligible", "qualification", "qualify",
        "turnover", "experience", "joint venture", "consortium", "blacklisted",
        "debarred",
    ),
    "submission_requirements": (
        "mandatory", "shall submit", "must submit", "certificate", "affidavit",
        "emd", "earnest money", "performance security", "bank guarantee",
        "technical bid", "financial bid", "price bid", "price schedule",
        "tender fee", "processing fee", "power of attorney", "undertaking",
        "gst registration", "experience certificate",
    ),
    "deadlines": (
        "deadline", "closing date", "due date", "submission date", "submit by",
        "last date", "bid submission", "submission end", "bid end date",
        "tender closing", "bid validity", "pre-bid meeting",
    ),
    "tender_economics": (
        "estimated cost", "estimated tender cost", "estimated project cost",
        "tender value", "estimated value", "bid amount", "contract value",
        "project cost", "earnest money", "emd", "tender fee",
        "processing fee", "security deposit", "performance guarantee",
        "payment", "gst", "tax", "price schedule",
        "price escalation", "purchase preference", "preferred over",
    ),
    "commercial_risks": (
        "penalty", "liquidated damages", "termination", "liability", "warranty",
        "indemnity", "arbitration", "disqualified", "disqualification",
    ),
}

_NEGATED_EXCLUSION = re.compile(
    r"\bnot\b(?:\W+\w+){0,5}\W+\b(?:blacklisted|debarred|disqualified)\b",
    re.IGNORECASE,
)
_EXPLICIT_EXCLUSION = re.compile(
    r"\b(?:ineligible|not eligible|not permitted to bid|"
    r"(?:shall|will|may|can)\s+be\s+disqualified|"
    r"liable\s+to\s+be\s+disqualified|"
    r"(?:shall|will|may)\s+be\s+blacklisted|"
    r"(?:shall|will|may)\s+be\s+debarred)\b",
    re.IGNORECASE,
)
_CURRENCY_AMOUNT = re.compile(
    r"(?P<currency>INR|Rs\.?|₹|\$|€|£|USD|EUR|GBP|JPY|AUD|CAD|AED)[\s():=]*"
    r"(?P<amount>(?:\d{1,3}(?:,\d{2,3})+|\d+)(?:\.\d+)?)",
    re.IGNORECASE,
)
_AMOUNT_PATTERN = r"(?:\d{1,3}(?:,\d{2,3})+|\d+)(?:\.\d+)?"
_TENDER_VALUE_LABEL = re.compile(
    r"\b(?:estimated\s+(?:tender\s+)?(?:cost|value|amount)|"
    r"(?:tender|contract|project)\s+(?:value|cost|amount)|"
    r"estimated\s+project\s+cost|bid\s+value|"
    r"cost\s+of\s+(?:the\s+)?(?:tender|contract|project|work))\b",
    re.IGNORECASE,
)
_DEADLINE_LABEL = re.compile(
    r"\b(?:deadline|closing date|due date|submission date|submit by|last date|"
    r"bid submission(?:\s+end(?:\s+date)?)?|submission end(?: date)?|"
    r"bid end date|tender closing|pre-bid meeting)\b",
    re.IGNORECASE,
)
_LABELED_TENDER_AMOUNT = re.compile(
    rf"\b(?:estimated\s+(?:tender\s+)?(?:cost|value|amount)|"
    rf"(?:tender|contract|project)\s+(?:value|cost|amount)|"
    rf"estimated\s+project\s+cost|bid\s+value|"
    rf"cost\s+of\s+(?:the\s+)?(?:tender|contract|project|work))\b"
    rf"[^0-9₹$€£]{{0,30}}(?P<amount>{_AMOUNT_PATTERN})",
    re.IGNORECASE,
)
_CURRENCY_SUFFIX_AMOUNT = re.compile(
    rf"(?P<amount>{_AMOUNT_PATTERN})\s*(?P<currency>INR|Rs\.?|USD|EUR|GBP|JPY|AUD|CAD|AED)",
    re.IGNORECASE,
)
_DATE_VALUE = re.compile(
    r"\b(?:"
    r"\d{4}-\d{1,2}-\d{1,2}"
    r"|\d{1,2}[./-]\d{1,2}[./-]\d{2,4}"
    r"|\d{1,2}\s+(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?)\s+\d{4}"
    r"|(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{4}"
    r")\b"
    r"(?:\s*(?:(?:at|,|upto|up to)\s*)?\d{1,2}:\d{2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?)?)?",
    re.IGNORECASE,
)
_DATE_ONLY_VALUE = re.compile(
    r"\b(?:"
    r"\d{4}-\d{1,2}-\d{1,2}"
    r"|\d{1,2}[./-]\d{1,2}[./-]\d{2,4}"
    r"|\d{1,2}\s+(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?)\s+\d{4}"
    r"|(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{4}"
    r")\b",
    re.IGNORECASE,
)
_TIMELINE_EVENT = re.compile(
    r"\b(?:invitation|inviting|publication|published|pre[- ]bid|query|"
    r"clarification|download|availability|submission|bid opening|opening of|"
    r"technical bid|financial bid|award|completion|site visit|inspection|"
    r"conference|meeting|last date|closing|sale of tender|tender schedule)\b",
    re.IGNORECASE,
)
_IMPORTANT_TIMELINE_EVENTS = (
    ("invitation", re.compile(r"\b(?:invitation|inviting|publication of e?[- ]?nit|nit publication|issue of e?[- ]?nit)\b", re.IGNORECASE)),
    ("pre_bid", re.compile(r"\b(?:pre[- ]?bid|pre[- ]?tender)\b", re.IGNORECASE)),
    ("document_download", re.compile(r"\b(?:download(?:ing)?|sale|availability)\b.*\b(?:tender|document|bid)\b|\b(?:tender|document|bid)\b.*\b(?:download(?:ing)?|sale|availability)\b", re.IGNORECASE)),
    ("bid_submission", re.compile(r"\b(?:submission|submit|last date|deadline|closing)\b.*\b(?:tenders?|bids?|proposals?|offers?)\b|\b(?:tenders?|bids?|proposals?|offers?)\b.*\b(?:submission|submit|last date|deadline|closing)\b", re.IGNORECASE)),
    ("technical_bid_opening", re.compile(r"\b(?:technical|tech)\b.*\b(?:bid|tender)\b.*\bopening\b|\bopening\b.*\b(?:technical|tech)\b.*\b(?:bid|tender)\b", re.IGNORECASE)),
    ("financial_bid_opening", re.compile(r"\b(?:financial|price)\b.*\b(?:bid|tender)\b.*\bopening\b|\bopening\b.*\b(?:financial|price)\b.*\b(?:bid|tender)\b", re.IGNORECASE)),
)
_TIMELINE_TIME = re.compile(
    r"\b\d{1,2}:\d{2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?)?\b",
    re.IGNORECASE,
)
_TIMELINE_HEADING = re.compile(
    r"\b(?:time schedule|tender schedule|schedule of (?:tender )?activities|"
    r"schedule of events|important dates|key dates|critical dates|"
    r"date and time schedule|tender programme)\b",
    re.IGNORECASE,
)
_TIMELINE_SECTION_END = re.compile(
    r"^\s*\d{1,2}\.\s+(?:ELIGIBILITY|PRE[- ]?QUALIFICATION|"
    r"QUALIFICATION|SCOPE OF WORK|INSTRUCTIONS|TERMS AND CONDITIONS|"
    r"FINANCIAL BID|TECHNICAL SPECIFICATION)\b",
    re.IGNORECASE,
)
_PREQUAL_SECTION_END = re.compile(
    r"^\s*\d{1,2}\.\s+(?:SCOPE OF WORK|INSTRUCTIONS|TERMS AND CONDITIONS|"
    r"FINANCIAL BID|TECHNICAL SPECIFICATION|BID SECURITY|SUBMISSION OF BIDS)\b",
    re.IGNORECASE,
)
_PREQUAL_HEADING = re.compile(
    r"\b(?:pre[- ]?qualification criteria|prequalification requirements|"
    r"technical qualification criteria|eligibility criteria)\b",
    re.IGNORECASE,
)
_PREQUAL_NUMBERED = re.compile(r"^\s*(\d{1,3}\.\d+(?:\.\d+)*)\s*(.*)$")
_PREQUAL_NUMBER_ONLY = re.compile(r"^\s*(\d{1,3}\.\d+(?:\.\d+)*)\s*$")
_COMPLETION_LABEL = re.compile(
    r"\b(?:completion period|period of completion|time for completion|"
    r"completion time|contract duration|execution period)\b",
    re.IGNORECASE,
)
_DURATION_VALUE = re.compile(
    r"\b\d+(?:\.\d+)?\s*(?:days?|weeks?|months?|years?)\b",
    re.IGNORECASE,
)
_FINANCIAL_LABELS = (
    (
        "Tender value",
        re.compile(
            r"\b(?:tender|contract|project)\s+value\b",
            re.IGNORECASE,
        ),
    ),
    (
        "Estimated cost",
        re.compile(
            r"\b(?:estimated\s+(?:(?:tender|project)\s+)?(?:cost|value|amount)|"
            r"(?:tender|contract|project)\s+cost|"
            r"cost\s+of\s+(?:the\s+)?(?:tender|contract|project|work))\b",
            re.IGNORECASE,
        ),
    ),
    (
        "Earnest money deposit (EMD)",
        re.compile(r"\b(?:earnest money|emd)\b", re.IGNORECASE),
    ),
    (
        "Bid security",
        re.compile(r"\bbid security\b", re.IGNORECASE),
    ),
    ("Tender fee", re.compile(r"\btender fees?\b", re.IGNORECASE)),
    ("Processing fee", re.compile(r"\bprocessing fees?\b", re.IGNORECASE)),
    (
        "Performance security",
        re.compile(r"\b(?:performance security|performance guarantee)\b", re.IGNORECASE),
    ),
    ("Security deposit", re.compile(r"\bsecurity deposit\b", re.IGNORECASE)),
)


def _financial_cost_type(text: str) -> str | None:
    for label, pattern in _FINANCIAL_LABELS:
        if pattern.search(text):
            return label
    return None


def _layout_amount_cost_types(
    page: dict[str, Any],
) -> dict[tuple[str, int], str]:
    rows = page.get("layout_rows", [])
    labels: list[tuple[float, float, str]] = []
    amount_cells: list[tuple[float, float, str]] = []
    for row in rows:
        y = row["y"]
        for cell in row["cells"]:
            text = cell["text"]
            label = _financial_cost_type(text)
            if label:
                labels.append((y, cell["x"], label))
            matches = list(_CURRENCY_AMOUNT.finditer(text))
            if not matches:
                matches = list(_CURRENCY_SUFFIX_AMOUNT.finditer(text))
            for match in matches:
                currency = match.groupdict().get("currency")
                amount = match.group("amount")
                value = f"{currency} {amount}" if currency else amount
                amount_cells.append((y, cell["x"], value))

    occurrence_by_value: dict[str, int] = {}
    result: dict[tuple[str, int], str] = {}
    for y, x, value in amount_cells:
        occurrence = occurrence_by_value.get(value, 0)
        occurrence_by_value[value] = occurrence + 1
        nearby_labels = [
            (abs(label_x - x) + (label_y - y) * 0.15, label)
            for label_y, label_x, label in labels
            if 0 <= label_y - y <= 240
        ]
        if nearby_labels:
            result[(value, occurrence)] = min(nearby_labels, key=lambda item: item[0])[1]
    return result


def _extract_pdf(filename: str, content: bytes) -> list[dict[str, Any]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
    except (PdfReadError, OSError, ValueError) as exc:
        raise ValueError(f"{filename}: unable to read this PDF") from exc
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"{filename}: PDF exceeds the {MAX_PDF_PAGES}-page analysis limit")
    pages = []
    total_chars = 0
    for page_number, page in enumerate(reader.pages, start=1):
        positioned_fragments: list[tuple[float, float, str]] = []

        def capture_positioned_text(
            fragment: str,
            current_matrix: Any,
            text_matrix: Any,
            font: Any,
            font_size: float,
        ) -> None:
            del current_matrix, font, font_size
            fragment = " ".join(fragment.split())
            if fragment:
                positioned_fragments.append((
                    float(text_matrix[5]),
                    float(text_matrix[4]),
                    fragment,
                ))

        text = (page.extract_text(visitor_text=capture_positioned_text) or "").strip()
        total_chars += len(text)
        if total_chars > MAX_EXTRACTED_CHARS_PER_DOCUMENT:
            remaining = max(0, MAX_EXTRACTED_CHARS_PER_DOCUMENT - (total_chars - len(text)))
            text = text[:remaining]
        positioned_lines: list[tuple[float, list[tuple[float, str]]]] = []
        for y, x, fragment in sorted(positioned_fragments, key=lambda item: (-item[0], item[1])):
            if not positioned_lines or abs(positioned_lines[-1][0] - y) > 2:
                positioned_lines.append((y, [(x, fragment)]))
            else:
                positioned_lines[-1][1].append((x, fragment))
        layout_text = "\n".join(
            "  ".join(fragment for _, fragment in sorted(fragments))
            for _, fragments in positioned_lines
        )
        layout_rows = [
            {
                "y": y,
                "cells": [
                    {"x": x, "text": fragment}
                    for x, fragment in sorted(fragments)
                ],
            }
            for y, fragments in positioned_lines
        ]
        pages.append({
            "filename": filename,
            "location": f"page {page_number}",
            "text": text,
            "layout_text": layout_text[:MAX_EXTRACTED_CHARS_PER_DOCUMENT],
            "layout_rows": layout_rows,
        })
        if total_chars >= MAX_EXTRACTED_CHARS_PER_DOCUMENT:
            break
    if not any(page["text"] for page in pages):
        raise ValueError(f"{filename}: no searchable text found; scanned PDFs need OCR before analysis")
    return pages


def _extract_docx(filename: str, content: bytes) -> list[dict[str, Any]]:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError(f"{filename}: expanded DOCX content exceeds the 30 MB limit")
        document = Document(io.BytesIO(content))
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(f"{filename}: unable to read this DOCX file") from exc
    paragraphs = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            paragraphs.append({"filename": filename, "location": f"paragraph {number}", "text": text})
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append({
                    "filename": filename,
                    "location": f"table {table_number}, row {row_number}",
                    "text": " | ".join(cells),
                })
    if not paragraphs:
        raise ValueError(f"{filename}: no readable paragraphs or table rows found")
    return paragraphs


def _extract_text(filename: str, content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{filename}: text file must be UTF-8 encoded") from exc
    if not text.strip():
        raise ValueError(f"{filename}: text file is empty")
    paragraphs = []
    total_chars = 0
    for number, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        remaining = MAX_EXTRACTED_CHARS_PER_DOCUMENT - total_chars
        if remaining <= 0:
            break
        line = line[:remaining]
        total_chars += len(line)
        paragraphs.append({"filename": filename, "location": f"line {number}", "text": line})
    return paragraphs


@traced_span("rpg_foresight.tool.extract_document_text")
def extract_document_pages(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Extract text with locations; never persists the uploaded source file."""
    if len(content) > MAX_DOCUMENT_SIZE:
        raise ValueError(f"{filename}: document exceeds the 10 MB limit")
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == ".pdf":
        locations = _extract_pdf(filename, content)
    elif extension == ".docx":
        locations = _extract_docx(filename, content)
    elif extension == ".txt":
        locations = _extract_text(filename, content)
    else:
        raise ValueError(f"{filename}: supported formats are PDF, DOCX, and TXT")
    from opentelemetry import trace

    span = trace.get_current_span()
    span.set_attribute("rpg_foresight.document.location_count", len(locations))
    span.set_attribute("rpg_foresight.document.format", extension.lstrip("."))
    return locations


def _timeline_event_type(event: str) -> str | None:
    for event_type, pattern in _IMPORTANT_TIMELINE_EVENTS:
        if pattern.search(event):
            return event_type
    return None


def _normalize_timeline_value(value: str) -> tuple[str, datetime | None]:
    date_match = _DATE_ONLY_VALUE.search(value)
    if date_match is None:
        return value, None

    raw_date = date_match.group(0)
    date_formats = (
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d %B %Y",
        "%d %b %Y",
        "%B %d, %Y",
        "%b %d, %Y",
    )
    parsed_date = None
    for date_format in date_formats:
        try:
            parsed_date = datetime.strptime(raw_date, date_format)
            break
        except ValueError:
            continue
    if parsed_date is None:
        return value, None

    time_match = _TIMELINE_TIME.search(value[date_match.end():])
    parsed_time = None
    display_time = ""
    if time_match:
        raw_time = time_match.group(0).strip()
        for time_format in ("%I:%M:%S %p", "%I:%M %p", "%H:%M:%S", "%H:%M"):
            try:
                parsed_time = datetime.strptime(raw_time, time_format)
                break
            except ValueError:
                continue
        if parsed_time:
            display_time = parsed_time.strftime("%I:%M %p")
        else:
            display_time = raw_time

    display_value = parsed_date.strftime("%d %b %Y")
    if display_time:
        display_value = f"{display_value} · {display_time}"
    sort_value = parsed_date.replace(
        hour=parsed_time.hour if parsed_time else 0,
        minute=parsed_time.minute if parsed_time else 0,
        second=parsed_time.second if parsed_time else 0,
    )
    return display_value, sort_value


@traced_span("rpg_foresight.tool.extract_document_evidence")
def find_document_evidence(pages: list[dict[str, Any]]) -> dict[str, Any]:
    """Return keyword-grounded findings and excerpts with source locations."""
    findings: dict[str, list[dict[str, str]]] = {category: [] for category in _FINDING_TERMS}
    disqualifiers: list[dict[str, str]] = []
    extracted_financial_amounts: list[dict[str, str]] = []
    extracted_tender_values: list[dict[str, str]] = []
    extracted_deadlines: list[dict[str, str]] = []
    extracted_timeline: list[dict[str, str]] = []
    extracted_tender_periods: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    seen_amount_occurrences: dict[tuple[str, str, str], int] = {}
    seen_tender_values: set[tuple[str, str, str]] = set()
    seen_deadlines: set[tuple[str, str, str]] = set()
    seen_timeline: set[tuple[str, str, str, str]] = set()
    seen_periods: set[tuple[str, str, str]] = set()

    def clauses(text: str) -> list[str]:
        return [
            " ".join(raw.split())
            for raw in re.split(
                r"(?<!Rs\.)(?<=[.!?])\s+|\n+",
                text,
                flags=re.IGNORECASE,
            )
            if raw.strip()
        ]

    for page in pages:
        text = page["text"]
        layout_text = page.get("layout_text", "")
        source_texts = [layout_text, text] if layout_text else [text]
        layout_cost_types = _layout_amount_cost_types(page)
        for sentence in clauses(text):
            for category, keywords in _FINDING_TERMS.items():
                if len(findings[category]) >= 12:
                    continue
                lowered = sentence.casefold()
                if not any(keyword in lowered for keyword in keywords):
                    continue
                key = (category, page["filename"], sentence.casefold())
                if key in seen:
                    continue
                seen.add(key)
                item = {
                    "document": page["filename"],
                    "location": page["location"],
                    "excerpt": sentence[:500],
                }
                findings[category].append(item)
                if (
                    category in {"eligibility", "commercial_risks"}
                    and not _NEGATED_EXCLUSION.search(sentence)
                    and _EXPLICIT_EXCLUSION.search(sentence)
                ):
                    disqualifiers.append(item)

        for source_text in source_texts:
            amount_occurrences_on_page: dict[str, int] = {}
            for sentence in clauses(source_text):
                citation = {
                    "document": page["filename"],
                    "location": page["location"],
                    "excerpt": sentence[:2000],
                }
                amount_matches = list(_CURRENCY_AMOUNT.finditer(sentence))
                if not amount_matches:
                    amount_matches = list(_CURRENCY_SUFFIX_AMOUNT.finditer(sentence))
                for match in amount_matches:
                    currency = match.groupdict().get("currency")
                    amount = match.group("amount")
                    value = f"{currency} {amount}" if currency else amount
                    occurrence_key = (page["filename"], page["location"], value)
                    occurrence = amount_occurrences_on_page.get(value, 0)
                    amount_occurrences_on_page[value] = occurrence + 1
                    if occurrence < seen_amount_occurrences.get(occurrence_key, 0):
                        continue
                    seen_amount_occurrences[occurrence_key] = occurrence + 1
                    amount_item = {"value": value, **citation}
                    cost_type = (
                        _financial_cost_type(sentence)
                        or layout_cost_types.get((value, occurrence))
                    )
                    if cost_type:
                        amount_item["cost_type"] = cost_type
                    extracted_financial_amounts.append(amount_item)

                value_label = _TENDER_VALUE_LABEL.search(sentence)
                if value_label:
                    value_match = next(
                        (match for match in amount_matches if match.start() >= value_label.end()),
                        None,
                    )
                    if value_match:
                        currency = value_match.groupdict().get("currency")
                        amount = value_match.group("amount")
                        value = f"{currency} {amount}" if currency else amount
                        key = (page["filename"], page["location"], value)
                        if key not in seen_tender_values:
                            extracted_tender_values.append({"value": value, **citation})
                            seen_tender_values.add(key)

                for date_match in _DATE_VALUE.finditer(sentence):
                    value = date_match.group(0).strip(" ,;:")
                    key = (page["filename"], page["location"], value)
                    if _DEADLINE_LABEL.search(sentence) and key not in seen_deadlines:
                        extracted_deadlines.append({"value": value, **citation})
                        seen_deadlines.add(key)
                if _COMPLETION_LABEL.search(sentence):
                    for duration_match in _DURATION_VALUE.finditer(sentence):
                        value = duration_match.group(0)
                        key = (page["filename"], page["location"], value)
                        if key not in seen_periods:
                            extracted_tender_periods.append({"value": value, **citation})
                            seen_periods.add(key)

        # Column-wise PDF extraction can put a table label and its value on
        # adjacent lines. Only pair a label with an amount in the same line or
        # immediately following currency/value-only lines.
        page_text = " ".join(text.split())
        page_citation = {
            "document": page["filename"],
            "location": page["location"],
            "excerpt": page_text[:2000],
        }
        tender_lines = text.splitlines()
        for line_index, line in enumerate(tender_lines):
            for label in _TENDER_VALUE_LABEL.finditer(line):
                amount_match = _CURRENCY_AMOUNT.search(line[label.end():])
                matched_text = line
                if amount_match is None:
                    following_lines: list[str] = []
                    for following_line in tender_lines[line_index + 1:line_index + 4]:
                        candidate = following_line.strip()
                        if not candidate:
                            continue
                        if re.fullmatch(r"[\[(]?(?:INR|Rs\.?|₹)[\])]?", candidate, re.IGNORECASE):
                            following_lines.append(candidate)
                            continue
                        if re.fullmatch(_AMOUNT_PATTERN + r"\s*/?-?", candidate):
                            following_lines.append(candidate)
                            matched_text = " ".join((line, *following_lines))
                            amount_match = _CURRENCY_AMOUNT.search(matched_text[label.end():])
                            if amount_match is None:
                                amount_match = _LABELED_TENDER_AMOUNT.search(matched_text)
                            break
                        break
                if amount_match:
                    currency = amount_match.groupdict().get("currency")
                    amount = amount_match.group("amount")
                    value = f"{currency} {amount}" if currency else amount
                    key = (page["filename"], page["location"], value)
                    if key not in seen_tender_values:
                        extracted_tender_values.append({
                            "value": value,
                            "document": page["filename"],
                            "location": page["location"],
                            "excerpt": " ".join(matched_text.split())[:500],
                        })
                        seen_tender_values.add(key)
                    if not any(
                        item["document"] == page["filename"]
                        and item["location"] == page["location"]
                        and item["value"] == value
                        for item in extracted_financial_amounts
                    ):
                        extracted_financial_amounts.append({
                            "value": value,
                            "cost_type": (
                                _financial_cost_type(matched_text)
                                or "Tender value"
                            ),
                            **page_citation,
                        })
        for label in _DEADLINE_LABEL.finditer(page_text):
            context = page_text[label.end():label.end() + 180]
            date_match = _DATE_VALUE.search(context)
            if date_match:
                value = date_match.group(0).strip(" ,;:")
                key = (page["filename"], page["location"], value)
                if key not in seen_deadlines:
                    extracted_deadlines.append({"value": value, **page_citation})
                    seen_deadlines.add(key)
        if _COMPLETION_LABEL.search(page_text):
            for duration_match in _DURATION_VALUE.finditer(page_text):
                value = duration_match.group(0)
                key = (page["filename"], page["location"], value)
                if key not in seen_periods:
                    extracted_tender_periods.append({"value": value, **page_citation})
                    seen_periods.add(key)

        timeline_text = (
            layout_text
            if layout_text and _TIMELINE_HEADING.search(layout_text)
            else text
        )
        timeline_active = False
        for line in timeline_text.splitlines():
            normalized_line = " ".join(line.split())
            if not normalized_line:
                continue
            if _TIMELINE_HEADING.search(normalized_line):
                timeline_active = True
            elif timeline_active and _TIMELINE_SECTION_END.search(normalized_line):
                timeline_active = False
            for date_match in _DATE_ONLY_VALUE.finditer(normalized_line):
                event = normalized_line[:date_match.start()]
                event = re.sub(r"^\s*\d+(?:[.)]\s*|\s+)", "", event)
                event = event.strip(" \t:;,.—-")
                event_type = _timeline_event_type(event)
                if event_type is None or (
                    not timeline_active and not _TIMELINE_EVENT.search(event)
                ):
                    continue
                full_date_match = _DATE_VALUE.match(normalized_line[date_match.start():])
                raw_value = (
                    full_date_match.group(0).strip(" ,;:")
                    if full_date_match
                    else date_match.group(0).strip(" ,;:")
                )
                value, sort_value = _normalize_timeline_value(raw_value)
                key = (page["filename"], event_type, value)
                if key in seen_timeline:
                    continue
                extracted_timeline.append({
                    "event": event[:1000],
                    "event_type": event_type,
                    "value": value,
                    "sort_value": sort_value.isoformat() if sort_value else "",
                    "document": page["filename"],
                    "location": page["location"],
                    "excerpt": normalized_line[:2000],
                })
                seen_timeline.add(key)

    extracted_timeline.sort(
        key=lambda item: (
            item["sort_value"] or "9999",
            item["document"].casefold(),
            item["location"],
        ),
    )
    for item in extracted_timeline:
        item.pop("sort_value")
    result = {
        "documents_analyzed": len({page["filename"] for page in pages}),
        "text_locations_reviewed": len(pages),
        "findings": findings,
        "extracted_financial_amounts": extracted_financial_amounts,
        "extracted_tender_values": extracted_tender_values,
        "extracted_deadlines": extracted_deadlines,
        "extracted_timeline": extracted_timeline,
        "extracted_tender_periods": extracted_tender_periods,
        "prequalification_criteria": _collect_prequalification_criteria(pages),
        "explicit_disqualifier_evidence": disqualifiers,
        "limitations": [
            "Findings are keyword-based excerpts, not legal or eligibility determinations.",
            "PDF page references are provided; DOCX and TXT use paragraph and line references.",
            "Scanned images and OCR-only text may not be extracted; verify values and dates against the original document.",
        ],
    }
    from opentelemetry import trace

    span = trace.get_current_span()
    span.set_attribute("rpg_foresight.documents_analyzed", result["documents_analyzed"])
    span.set_attribute("rpg_foresight.text_locations_reviewed", result["text_locations_reviewed"])
    span.set_attribute(
        "rpg_foresight.finding_count",
        sum(len(items) for items in findings.values()),
    )
    span.set_attribute(
        "rpg_foresight.explicit_exclusion_count",
        len(disqualifiers),
    )
    return result


def _collect_prequalification_criteria(
    pages: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    criteria: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    in_section = False

    def add_line(line: str, page: dict[str, Any]) -> None:
        nonlocal current
        if current is None or not line:
            return
        if line not in current["_lines"]:
            current["_lines"].append(line)
        references = current["references"]
        if references and (
            references[-1]["document"] == page["filename"]
            and references[-1]["location"] == page["location"]
        ):
            if line not in references[-1]["excerpt"]:
                references[-1]["excerpt"] = f'{references[-1]["excerpt"]} {line}'[:4000]
        else:
            references.append({
                "document": page["filename"],
                "location": page["location"],
                "excerpt": line[:4000],
            })

    def finish_current() -> None:
        if current is None:
            return
        excerpt = " ".join(current.pop("_lines")).strip()
        if excerpt:
            current["excerpt"] = excerpt
            criteria.append(current.copy())

    for page in pages:
        if current is not None and current["references"] and (
            current["references"][-1]["document"] != page["filename"]
        ):
            finish_current()
            current = None
            in_section = False
        for raw_line in page["text"].splitlines():
            line = " ".join(raw_line.split())
            if not line:
                continue
            if _PREQUAL_HEADING.search(line):
                in_section = True
                remainder = _PREQUAL_HEADING.sub("", line).strip(" :-")
                if remainder:
                    add_line(remainder, page)
                continue
            if not in_section:
                continue
            if _PREQUAL_SECTION_END.search(line) or re.match(
                r"^\s*\d{1,2}\.\s+[A-Z][A-Z\s:&()/-]{4,}\s*$",
                line,
            ):
                finish_current()
                current = None
                in_section = False
                continue
            number_only = _PREQUAL_NUMBER_ONLY.match(line)
            numbered = _PREQUAL_NUMBERED.match(line)
            if number_only or numbered:
                finish_current()
                criterion_id = (number_only or numbered).group(1)
                criterion_text = numbered.group(2).strip() if numbered else ""
                current = {
                    "criterion_id": criterion_id,
                    "_lines": [],
                    "references": [],
                }
                add_line(criterion_text, page)
                continue
            if current is not None and not re.fullmatch(r"(?:page\s*)?\d+", line, re.IGNORECASE):
                add_line(line, page)
    finish_current()
    return criteria
