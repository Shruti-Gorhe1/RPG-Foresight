from __future__ import annotations

import io
import json
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from docx import Document
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from api.auth import CurrentUser, get_current_user
from api.routers.pursuit_copilot import router
from agents.pursuit.document_agent import (
    _summary_numbers_are_grounded,
    _apply_prequalification_summaries,
    document_agent,
    _prequalification_summary_inputs,
    _select_review_evidence,
    _validated_document_summaries,
    _validate_summary_item,
)
from agents.pursuit.model import parse_model_answer
from agents.pursuit.action_planner_agent import action_planner_agent
from agents.tools.pursuit_review_tools import (
    prepare_pursuit_action_plan,
    review_tender_risks,
)
from db.base import Base, get_db_session
from db.models import (
    PursuitAssessment,
    RoutingRule,
    SubsidiaryRoute,
    TenderEvent,
    TenderOutcomeSnapshot,
)
from services.pursuit_documents import extract_document_pages, find_document_evidence
from services.pursuit_chat import (
    answer_prerequisite_question,
    answer_explicit_tender_value_question,
    answer_tender_value_cost_comparison,
    collect_prerequisite_evidence,
    collect_tender_risk_evidence,
    retrieve_tender_evidence,
)


@pytest_asyncio.fixture
async def copilot_client():
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    app = FastAPI()
    app.include_router(router, prefix="/api/v1/pursuit")

    async def override_db():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_db_session] = override_db
    app.dependency_overrides[get_current_user] = lambda: CurrentUser("bid-owner", "TENANT-A")
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, session_factory, app
    await engine.dispose()


def test_tender_chat_retrieves_relevant_cited_chunks_only():
    pages = [
        {
            "filename": "eligibility.pdf",
            "location": "page 4",
            "text": "The bidder must have completed similar contracts worth INR 50 lakh.",
        },
        {
            "filename": "schedule.pdf",
            "location": "page 9",
            "text": "The online bid submission deadline is 30 December 2026.",
        },
    ]

    evidence = retrieve_tender_evidence("What is the bid submission deadline?", pages)

    assert evidence
    assert evidence[0]["document"] == "schedule.pdf"
    assert evidence[0]["location"] == "page 9"
    assert evidence[0]["source_id"] == "S1"
    assert retrieve_tender_evidence("How is the weather?", pages) == []


@pytest.mark.parametrize(
    ("response", "expected"),
    [
        ("- Verify eligibility.\n- Submit audited accounts.", "- Verify eligibility.\n- Submit audited accounts."),
        ('{"answer":"- Verify eligibility."}', "- Verify eligibility."),
        ('{"answer":"- Verify eligibility." ', "- Verify eligibility."),
        ('{"answer":"Verify\\n eligibility." ', "Verify\n eligibility."),
    ],
)
def test_model_answer_parser_accepts_plain_text_and_truncated_json(
    response: str,
    expected: str,
):
    assert parse_model_answer(response) == expected


def test_model_answer_parser_caps_overlong_answers_without_rejecting_them():
    response = "- requirement " * 120

    answer = parse_model_answer(response)

    assert len(answer.split()) <= 90
    assert answer.startswith("- requirement")


def test_prerequisite_question_returns_qualification_points_not_amounts_or_references():
    answer = answer_prerequisite_question(
        "What are the pre-requisites for the tender?",
        {
            "prequalification_criteria": [
                {
                    "criterion_id": "5.1",
                    "excerpt": (
                        "The applicant should have satisfactorily completed "
                        "similar works during the last seven years."
                    ),
                    "summary": "Completed similar works during the last seven years.",
                },
                {
                    "criterion_id": "5.2",
                    "excerpt": "Average annual financial turnover must be INR 12,00,000.",
                },
            ],
            "findings": {
                "eligibility": [
                    {
                        "excerpt": (
                            "The applicant should have satisfactorily completed "
                            "similar works during the last seven years."
                        ),
                    },
                    {
                        "excerpt": "Bidders must not be blacklisted.",
                    },
                ],
                "submission_requirements": [
                    {"excerpt": "Submit a valid GST certificate."},
                ],
            },
        },
    )

    assert answer is not None
    assert "Eligibility and qualification requirements:" in answer["answer"]
    assert "similar works during the last seven years" in answer["answer"]
    assert "average annual financial turnover must be inr 12,00,000" in answer["answer"].casefold()
    assert "must not be blacklisted" in answer["answer"].casefold()
    assert "Required documents or submissions:" in answer["answer"]
    assert "Submit a valid GST certificate" in answer["answer"]
    assert answer["answer"].count("similar works during the last seven years") == 1
    assert "TD.pdf" not in answer["answer"]
    assert "page" not in answer["answer"].casefold()
    assert answer["sources"] == []


def test_prerequisite_evidence_filters_unrelated_contract_schedule_and_ocr_text():
    evidence = collect_prerequisite_evidence(
        "What are the prerequisites for this tender?",
        {
            "prequalification_criteria": [
                {
                    "criterion_id": "5.1",
                    "summary_source": "document_excerpt",
                    "excerpt": (
                        "Overheads and profits: 15% (Fifteen percent); "
                        "Time allowed for performance guarantee: 15 days; "
                        "The bidder must have completed similar electrical "
                        "audit work during the last seven years; "
                        "Average annual turnover must be INR 12,00,000; "
                        "Payment: 70% after draft report and 30% after final report."
                    ),
                },
            ],
            "findings": {
                "eligibility": [
                    {"excerpt": "The bidder must have completed similar work."},
                ],
                "submission_requirements": [
                    {"excerpt": "Submit a signed declaration."},
                ],
            },
        },
    )

    assert evidence is not None
    content = " ".join(item["text"] for item in evidence)
    assert "completed similar electrical audit work" in content
    assert "Average annual turnover" in content
    assert "Overheads and profits" not in content
    assert "Payment:" not in content
    assert "signed declaration" not in content


def test_risk_question_collects_relevant_tender_clauses_for_synthesis():
    evidence = collect_tender_risk_evidence(
        "What are the key-risks for this tender?",
        {
            "findings": {
                "commercial_risks": [
                    {
                        "excerpt": "Liquidated damages of 0.5% per week for delayed delivery.",
                    },
                ],
                "deadlines": [{"excerpt": "The bid submission deadline is 30 December."}],
                "submission_requirements": [
                    {"excerpt": "Submit a signed declaration with the bid."},
                ],
            },
            "explicit_disqualifier_evidence": [
                {
                    "excerpt": "Bidders may be disqualified if the declaration is missing.",
                },
            ],
        },
    )

    assert evidence is not None
    assert len(evidence) == 4
    assert {item["category"] for item in evidence} == {
        "commercial_risks",
        "deadlines",
        "submission_requirements",
        "explicit_disqualification",
    }
    assert any("Liquidated damages" in item["text"] for item in evidence)


def test_risk_question_returns_no_evidence_when_no_risk_relevant_clauses_exist():
    evidence = collect_tender_risk_evidence(
        "What are the key risks for this tender?",
        {
            "findings": {"commercial_risks": []},
            "explicit_disqualifier_evidence": [],
        },
    )

    assert evidence == []


def test_tender_value_question_uses_directly_labeled_value_not_deposit_percentage():
    retrieved = retrieve_tender_evidence(
        "What is the tender value?",
        [
            {
                "filename": "TD.pdf",
                "location": "page 37",
                "text": (
                    "The security deposit shall be 2.5% of the tendered value "
                    "of the work."
                ),
            },
            {
                "filename": "TD.pdf",
                "location": "page 12",
                "text": "Estimated cost Rs. 12,00,000. Tender value Rs. 24,000.",
            },
        ],
    )
    assert [item["location"] for item in retrieved] == ["page 12"]

    source_pages = [
        {
            "text": "Security deposit of 2.5% of the tendered value of the work.",
            "document": "TD.pdf",
            "location": "page 37",
        },
        {
            "text": "Estimated cost Rs. 12,00,000. Tender value Rs. 24,000.",
            "document": "TD.pdf",
            "location": "page 12",
        },
    ]

    answer = answer_explicit_tender_value_question(
        "What is the tender value?",
        source_pages,
    )

    assert answer is not None
    assert "Rs. 24,000" in answer["answer"]
    assert "2.5%" not in answer["answer"]
    assert [source["location"] for source in answer["sources"]] == ["page 12"]


def test_tender_value_question_refuses_when_only_unlabeled_percentage_is_found():
    answer = answer_explicit_tender_value_question(
        "What is the tender value?",
        [{
            "text": "Security deposit of 2.5% of the tendered value of the work.",
            "document": "TD.pdf",
            "location": "page 37",
        }],
    )

    assert answer is not None
    assert "couldn't identify a directly labeled tender value" in answer["answer"]
    assert answer["sources"] == []


def test_tender_value_question_reports_explicit_estimated_cost_when_no_value_is_labeled():
    answer = answer_explicit_tender_value_question(
        "What is the tender value?",
        [
            {
                "filename": "TD.pdf",
                "location": "page 12",
                "text": "Estimated cost put to Rs. 12,00,000/-.",
            },
            {
                "filename": "TD.pdf",
                "location": "page 37",
                "text": "Security deposit is 2.5% of tendered value.",
            },
        ],
        [
            {
                "value": "Rs. 12,00,000",
                "cost_type": "Estimated cost",
                "document": "TD.pdf",
                "location": "page 12",
                "excerpt": "Estimated cost put to Rs. 12,00,000/-.",
            },
            {
                "value": "2.5%",
                "cost_type": "Security deposit",
                "document": "TD.pdf",
                "location": "page 37",
                "excerpt": "Security deposit is 2.5% of tendered value.",
            },
        ],
    )

    assert answer is not None
    assert "does not separately label a tender value" in answer["answer"]
    assert "estimated cost" in answer["answer"]
    assert "Rs. 12,00,000" in answer["answer"]
    assert "2.5%" not in answer["answer"]
    assert [source["location"] for source in answer["sources"]] == ["page 12"]


def test_tender_cost_question_uses_estimated_cost_not_unrelated_criteria_amount():
    answer = answer_explicit_tender_value_question(
        "So what is the tender cost?",
        [
            {
                "filename": "TD.pdf",
                "location": "page 9",
                "text": "Under parallel contracts, amount: Rs. 10.",
            },
            {
                "filename": "TD.pdf",
                "location": "page 14",
                "text": "One work similar costing not less than Rs. 9,60,000, i.e. 80% of cost of tender.",
            },
            {
                "filename": "TD.pdf",
                "location": "page 12",
                "text": "Estimated cost: Rs. 12,00,000.",
            },
        ],
        [{
            "value": "Rs. 12,00,000",
            "cost_type": "Estimated cost",
            "document": "TD.pdf",
            "location": "page 12",
            "excerpt": "Estimated cost: Rs. 12,00,000.",
        }],
    )

    assert answer is not None
    assert "Rs. 12,00,000" in answer["answer"]
    assert "Rs. 10" not in answer["answer"]
    assert "Rs. 9,60,000" not in answer["answer"]
    assert [source["location"] for source in answer["sources"]] == ["page 12"]


def test_tender_cost_question_can_fall_through_to_document_retrieval():
    question = "What is the cost of the tender?"
    pages = [{
        "filename": "TD.pdf",
        "location": "page 12",
        "text": "Schedule of Rates: Total quoted amount Rs. 12,00,000.",
    }]
    answer = answer_explicit_tender_value_question(
        question,
        pages,
        return_none_if_not_found=True,
    )

    assert answer is None
    evidence = retrieve_tender_evidence(question, pages)
    assert evidence
    assert "Rs. 12,00,000" in evidence[0]["text"]


def test_tender_value_cost_comparison_explains_terms_and_deduplicates_cost():
    source_pages = [
        {
            "filename": "TD.pdf",
            "location": "page 9",
            "text": (
                "Under parallel contracts, the Class-I local Supplier will be "
                "preferred over Class-II local Supplier; amount: Rs. 10."
            ),
        },
        {
            "filename": "TD.pdf",
            "location": "page 12",
            "text": "Estimated cost put to Rs. 12,00,000/- (Excluding GST but including labour).",
        },
        {
            "filename": "TD.pdf",
            "location": "page 27",
            "text": "Estimated cost of work: Rs. 12,00,000/- As per Tender.",
        },
    ]
    financial_amounts = [
        {
            "value": "Rs. 10",
            "cost_type": "Tender value",
            "document": "TD.pdf",
            "location": "page 9",
            "excerpt": source_pages[0]["text"],
        },
        {
            "value": "Rs. 12,00,000",
            "cost_type": "Estimated cost",
            "document": "TD.pdf",
            "location": "page 12",
            "excerpt": source_pages[1]["text"],
        },
        {
            "value": "Rs. 12,00,000",
            "cost_type": "Estimated cost",
            "document": "TD.pdf",
            "location": "page 27",
            "excerpt": source_pages[2]["text"],
        },
    ]

    answer = answer_tender_value_cost_comparison(
        "What is the difference between tender value and tender cost?",
        source_pages,
        financial_amounts,
    )

    assert answer is not None
    assert "tender cost" in answer["answer"].casefold()
    assert "tender value" in answer["answer"].casefold()
    assert "Rs. 12,00,000" in answer["answer"]
    assert "Rs. 10" not in answer["answer"]
    assert "no separate, directly stated tender-value amount" in answer["answer"].casefold()
    assert [source["location"] for source in answer["sources"]] == ["page 12", "page 27"]


def test_tender_value_question_extracts_currency_and_rejects_labeled_percentages():
    answer = answer_explicit_tender_value_question(
        "What is the tender value?",
        [
            {
                "document": "TD.pdf",
                "location": "page 8",
                "text": "Tender value: INR 10,25,000.",
            },
            {
                "document": "TD.pdf",
                "location": "page 9",
                "text": "Project value 2.5% of the estimated work.",
            },
            {
                "document": "TD.pdf",
                "location": "page 10",
                "text": "Estimated Tender Value\n(INR)\n2,10,00,00,000",
            },
        ],
    )

    assert answer is not None
    assert "INR 10,25,000" in answer["answer"]
    assert "INR 2,10,00,00,000" in answer["answer"]
    assert "2.5%" not in answer["answer"]
    assert [source["location"] for source in answer["sources"]] == [
        "page 8",
        "page 10",
    ]


@pytest.mark.asyncio
async def test_copilot_runs_coordinated_agents_persists_citations_and_requires_human_decision(
    copilot_client,
    monkeypatch: pytest.MonkeyPatch,
):
    from opentelemetry import trace
    from opentelemetry.sdk.trace import TracerProvider
    from opentelemetry.sdk.trace.export import SimpleSpanProcessor
    from opentelemetry.sdk.trace.export.in_memory_span_exporter import InMemorySpanExporter

    client, session_factory, app = copilot_client
    model_calls: list[str] = []

    class FakeChatModel:
        async def ainvoke(self, prompt: str):
            model_calls.append(prompt)
            prompt_evidence = json.loads(prompt.split("Evidence JSON:\n", 1)[1])
            if "TENDER_RISK_SYNTHESIS_JSON" in prompt:
                clauses = prompt_evidence["tender_risk_clauses"]
                assert any(
                    "Liquidated damages apply to late delivery" in item["text"]
                    for item in clauses
                )
                return SimpleNamespace(content=json.dumps({
                    "answer": (
                        "The main documented exposure is liquidated damages if "
                        "delivery is late. This tender text doesn't identify an "
                        "encryption or bid-printout risk."
                    ),
                }))
            if "DOCUMENT_REVIEW_JSON" in prompt:
                evidence = prompt_evidence["evidence_by_category"]
                criteria = prompt_evidence["prequalification_criteria"]
                action_by_category = {
                    "eligibility": "Confirm bidder eligibility against the stated experience requirement.",
                    "submissions": "Prepare the required bid documents and certificates.",
                    "deadlines": "Schedule submission before the stated deadline.",
                    "commercial_terms": "Review the stated delivery and payment risks.",
                }
                return SimpleNamespace(content=json.dumps({
                    "analysis": "The uploaded tender text was reviewed and summarized.",
                    "summaries": [
                        {
                            "evidence_id": item["evidence_id"],
                            "summary": (
                                "The document records this financial term."
                                if category == "economics"
                                else action_by_category.get(
                                    category,
                                    "Complete the requirement stated in the tender.",
                                )
                            ),
                            **({
                                "description": (
                                    "Prepare proof that the bidder meets the required "
                                    "five years of relevant experience."
                                    if category == "eligibility"
                                    else "Prepare the specified bid documents and certificates."
                                )
                            } if category != "economics" else {}),
                        }
                        for category, items in evidence.items()
                        for item in items
                    ],
                    "prequalification_summaries": [
                        {
                            "criterion_key": criterion["criterion_key"],
                            "summary": criterion["excerpt"],
                        }
                        for criterion in criteria
                    ],
                }))
            return SimpleNamespace(content="Evidence-based analysis from the configured model.")

    monkeypatch.setattr("shared.llm.get_llm", lambda **_: FakeChatModel())
    span_exporter = InMemorySpanExporter()
    tracer_provider = TracerProvider()
    tracer_provider.add_span_processor(SimpleSpanProcessor(span_exporter))
    monkeypatch.setattr(trace, "get_tracer", tracer_provider.get_tracer)
    now = datetime.now(timezone.utc)
    async with session_factory() as session:
        session.add(TenderEvent(
            tenant_id="TENANT-A",
            event_id="copilot-tender",
            sender_entity_id="TENANT-A",
            tender_ref="RFP-100",
            project_name="West region cable supply",
            source_revision="rev-1",
            prior_stage=None,
            stage="Bid Submitted",
            occurred_at=now,
            material_categories=["Power cables"],
            expected_need_window=None,
            existing_win_probability_percent=82,
            probability_updated_at=now,
            source_type="crm",
            data_provenance="authorized CRM",
            region="West",
            tender_value=1_250_000,
            currency="INR",
            closing_at=now + timedelta(days=40),
        ))
        rule = RoutingRule(
            rule_id="cable-region-rule",
            tenant_id="TENANT-A",
            name="Cable bid",
            enabled=True,
            trigger_stages=["Bid Submitted"],
            min_win_probability_percent=70,
            purpose="Bid evaluation",
            channel="email",
            selected_recipient_email="buyer@example.test",
            selected_recipient_emails=["buyer@example.test"],
            regions=["West"],
            tender_currency="INR",
            minimum_tender_value=100_000,
            maximum_tender_value=2_000_000,
            max_event_age_seconds=86400,
            max_probability_age_seconds=604800,
            version=1,
        )
        rule.routes.append(SubsidiaryRoute(
            subsidiary_entity_id="SUB-A",
            material_categories=["Power cables"],
        ))
        session.add(rule)
        session.add(TenderEvent(
            tenant_id="TENANT-A",
            event_id="copilot-tender",
            sender_entity_id="TENANT-A",
            tender_ref="RFP-100",
            project_name="West region cable supply - latest revision",
            source_revision="rev-2",
            prior_stage="Bid Submitted",
            stage="Bid Submitted",
            occurred_at=now + timedelta(minutes=1),
            ingested_at=now + timedelta(minutes=1),
            material_categories=["Power cables"],
            expected_need_window=None,
            existing_win_probability_percent=82,
            probability_updated_at=now,
            source_type="crm",
            data_provenance="authorized CRM",
            region="West",
            tender_value=1_250_000,
            currency="INR",
            closing_at=now + timedelta(days=40),
        ))
        session.add_all([
            TenderOutcomeSnapshot(
                tenant_id="TENANT-A",
                source_event_id=f"historical-{index}",
                snapshot_key="bid-submitted",
                stage="Bid Submitted",
                material_categories=["Power cables"],
                region="West",
                tender_value=1_250_000,
                currency="INR",
                outcome="won" if index < 8 else "lost",
                is_demo=False,
                source_type="crm",
            )
            for index in range(10)
        ])
        await session.commit()

    document = (
        b"Bidder eligibility: minimum experience of five years is required.\n"
        b"Estimated cost: INR 12,00,000.\n"
        b"Mandatory: submit a valid GST certificate and performance security.\n"
        b"Submission deadline is 30 December 2026.\n"
        b"Liquidated damages apply to late delivery.\n"
        b"Security deposit is 2.5% of the tendered value of the work.\n"
        b"Pre-qualification criteria:\n"
        b"5.1 Minimum experience of five years is required."
    )
    response = await client.post(
        "/api/v1/pursuit/copilot-tender/analyze",
        files=[("files", ("RFP.txt", document, "text/plain"))],
    )
    assert response.status_code == 201, response.text
    result = response.json()
    assert "_document_sources" not in result["report"]
    assert result["recommendation"] == "bid"
    assert result["report"]["tender"]["project_name"] == "West region cable supply - latest revision"
    assert result["status"] == "pending_review"
    assert result["decision"] is None
    assert result["report"]["human_approval_required"] is True
    assert result["report"]["recommendation"]["external_actions_taken"] == []
    assert result["report"]["win_probability"]["training_tender_count"] == 10
    assert result["report"]["win_probability"]["probability_percent"] > 70
    assert result["report"]["tender"]["tender_value"] == 1_250_000
    assert result["report"]["tender"]["currency"] == "INR"
    assert result["report"]["tender"]["data_provenance"] == "authorized CRM"
    assert result["report"]["tender"]["probability_updated_at"] is not None
    assert len(model_calls) == 1
    assert result["report"]["model"]["provider"]
    assert result["report"]["trace_id"]
    assert result["report"]["documents"]["agent_analysis"] == "The uploaded tender text was reviewed and summarized."
    assert result["report"]["documents"]["economics_summaries"][0]["summary"] == (
        "The document records this financial term; amount: INR 12,00,000."
    )
    assert result["report"]["documents"]["processing_guidelines"]
    assert result["report"]["documents"]["prequalification_criteria"][0]["summary"] == (
        "Minimum experience of five years is required."
    )
    assert result["report"]["documents"]["prequalification_criteria"][0]["summary_source"] == "model"
    assert [item["agent"] for item in result["report"]["agent_trace"]] == [
        "pursuit_supervisor",
        "qualification_agent", "document_agent", "risk_agent", "action_planner_agent"
    ]
    assert len(result["report"]["recommendation"]["next_steps"]) == 3
    assert "Confirm bidder eligibility" in result["report"]["recommendation"]["next_steps"][0]["task"]
    assert all(
        step["description"].strip()
        for step in result["report"]["recommendation"]["next_steps"]
    )
    assert {
        item["parent_agent"] for item in result["report"]["tool_trace"]
    } == {
        "qualification_agent",
        "document_agent",
        "risk_agent",
        "action_planner_agent",
    }
    citations = result["report"]["documents"]["findings"]["eligibility"]
    assert citations[0]["document"] == "RFP.txt"
    assert citations[0]["location"] == "line 1"
    assert "five years" in citations[0]["excerpt"]

    chat_model_calls: list[str] = []

    async def grounded_chat_model(*args, **kwargs):
        chat_model_calls.append(args[0])
        if args[0] == "tender_prerequisite_synthesis":
            clauses = args[2]["qualification_clauses"]
            assert any("minimum experience of five years" in item["text"].casefold()
                       for item in clauses)
            assert not any("security deposit" in item["text"].casefold()
                           for item in clauses)
            return (
                '{"answer":"- Minimum five years of relevant electrical safety audit '
                'experience.\\n- Average annual turnover of at least INR 12,00,000." '
            )
        if args[0] == "tender_risk_synthesis":
            return json.dumps({
                "answer": (
                    "Late delivery may trigger liquidated damages."
                ),
            })
        return json.dumps({
            "answer": "The bidder must have at least five years of experience.",
            "source_ids": ["S1"],
        })

    monkeypatch.setattr(
        "api.routers.pursuit_copilot.reason_with_live_model",
        grounded_chat_model,
    )
    chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "What experience is required?", "history": []},
    )
    assert chat.status_code == 200, chat.text
    assert chat.json()["answer"] == "The bidder must have at least five years of experience."
    assert chat.json()["sources"] == []

    model_call_count = len(model_calls)
    tender_value_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "What is the tender value?", "history": []},
    )
    assert tender_value_chat.status_code == 200
    assert "does not separately label a tender value" in (
        tender_value_chat.json()["answer"]
    )
    assert "INR 12,00,000" in tender_value_chat.json()["answer"]
    assert tender_value_chat.json()["sources"] == []
    assert "2.5%" not in tender_value_chat.json()["answer"]
    assert len(model_calls) == model_call_count

    tender_cost_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "So what is the tender cost?", "history": []},
    )
    assert tender_cost_chat.status_code == 200
    assert "INR 12,00,000" in tender_cost_chat.json()["answer"]
    assert tender_cost_chat.json()["sources"] == []
    assert len(model_calls) == model_call_count

    prerequisite_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "What are the pre-requisites for the tender?", "history": []},
    )
    assert prerequisite_chat.status_code == 200
    assert prerequisite_chat.json()["answer"] == (
        "- Minimum five years of relevant electrical safety audit experience.\n"
        "- Average annual turnover of at least INR 12,00,000."
    )
    assert len(prerequisite_chat.json()["answer"].split()) <= 90
    assert prerequisite_chat.json()["sources"] == []
    assert len(model_calls) == model_call_count

    risk_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "What are the key-risks for this tender?", "history": []},
    )
    assert risk_chat.status_code == 200
    assert risk_chat.json()["answer"] == "Late delivery may trigger liquidated damages."
    assert "encryption" not in risk_chat.json()["answer"].casefold()
    assert risk_chat.json()["sources"] == []
    assert chat_model_calls == [
        "tender_document_chat",
        "tender_prerequisite_synthesis",
        "tender_risk_synthesis",
    ]

    comparison_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={
            "question": "What is the difference between tender value and tender cost?",
            "history": [],
        },
    )
    assert comparison_chat.status_code == 200
    assert "INR 12,00,000" in comparison_chat.json()["answer"]
    assert "no separate, directly stated tender-value amount" in (
        comparison_chat.json()["answer"].casefold()
    )
    assert comparison_chat.json()["sources"] == []
    assert len(model_calls) == model_call_count

    history = await client.get("/api/v1/pursuit/copilot-tender/assessments")
    assert history.status_code == 200
    assert len(history.json()) == 1

    decision = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/decision",
        json={"decision": "bid", "note": "Bid team reviewed cited eligibility requirements."},
    )
    assert decision.status_code == 200, decision.text
    assert decision.json()["status"] == "bid_approved"
    assert decision.json()["decision_by"] == "bid-owner"
    assert decision.json()["decision_note"].startswith("Bid team")

    duplicate_decision = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/decision",
        json={"decision": "no_bid", "note": "Changed decision"},
    )
    assert duplicate_decision.status_code == 409

    app.dependency_overrides[get_current_user] = lambda: CurrentUser("other-user", "TENANT-B")
    cross_tenant = await client.get(f"/api/v1/pursuit/assessment/{result['assessment_id']}")
    assert cross_tenant.status_code == 404
    cross_tenant_chat = await client.post(
        f"/api/v1/pursuit/assessment/{result['assessment_id']}/chat",
        json={"question": "What experience is required?"},
    )
    assert cross_tenant_chat.status_code == 404

    async with session_factory() as session:
        stored = await session.execute(select(PursuitAssessment))
        assessment = stored.scalar_one()
        assert assessment.document_names == ["RFP.txt"]
        assert "Bidder eligibility" in assessment.report["documents"]["findings"]["eligibility"][0]["excerpt"]
        assert "actual full uploaded file" not in str(assessment.report)

    next_steps = result["report"]["recommendation"]["next_steps"]
    tasks = [step["task"] for step in next_steps]
    assert len(tasks) == len(set(tasks))
    eligibility_step = next(step for step in next_steps if "eligibility" in step["task"])
    assert any(reference["document"] == "RFP.txt" for reference in eligibility_step["references"])
    assert "five years" in eligibility_step["description"]
    assert all("excerpt" not in reference for reference in eligibility_step["references"])
    guidelines = result["report"]["documents"]["processing_guidelines"]
    assert all(item["references"] for item in guidelines)
    assert all(
        "page" in reference["location"] or "line" in reference["location"]
        for item in guidelines
        for reference in item["references"]
    )

    spans = span_exporter.get_finished_spans()
    trace_root = next(span for span in spans if span.name == "rpg_foresight.pursuit.analyze")
    assert result["report"]["trace_id"] == f"{trace_root.context.trace_id:032x}"
    trace_spans = [span for span in spans if span.context.trace_id == trace_root.context.trace_id]
    span_by_id = {span.context.span_id: span for span in trace_spans}
    for agent, tool in (
        ("rpg_foresight.agent.qualification", "rpg_foresight.tool.match_tender_to_rules"),
        ("rpg_foresight.agent.document", "rpg_foresight.tool.extract_document_text"),
        ("rpg_foresight.agent.document", "rpg_foresight.tool.extract_document_evidence"),
        ("rpg_foresight.agent.risk", "rpg_foresight.tool.review_tender_risks"),
        ("rpg_foresight.agent.action_planner", "rpg_foresight.tool.prepare_pursuit_action_plan"),
    ):
        agent_span = next(span for span in trace_spans if span.name == agent)
        tool_span = next(span for span in trace_spans if span.name == tool)
        assert tool_span.parent is not None
        actual_parent = span_by_id[tool_span.parent.span_id]
        assert actual_parent.context.span_id == agent_span.context.span_id, (
            f"{tool} parent was {actual_parent.name}; spans: "
            + ", ".join(
                f"{span.name}<-{span_by_id[span.parent.span_id].name if span.parent and span.parent.span_id in span_by_id else 'root'}"
                for span in trace_spans
            )
        )
        assert tool_span.attributes["openinference.span.kind"] == "TOOL"
    llm_spans = [
        span for span in trace_spans
        if span.name == "rpg_foresight.llm.agent_reasoning"
    ]
    assert len(llm_spans) == 1
    assert all(span.attributes["openinference.span.kind"] == "LLM" for span in llm_spans)
    assert all(span.attributes.get("llm.model_name") for span in llm_spans)
    agent_span_ids = {
        span.context.span_id for span in trace_spans
        if span.attributes.get("openinference.span.kind") == "AGENT"
    }
    assert all(span.parent and span.parent.span_id in agent_span_ids for span in llm_spans)
    tracer_provider.shutdown()


def test_economic_summary_retains_amount_detected_in_the_document():
    summary = _validate_summary_item(
        {
            "evidence_id": "ECONOMICS-01",
            "summary": "The tender requires an earnest money deposit.",
        },
        "economics",
        {
            "ECONOMICS-01": {
                "category": "economics",
                "document": "Tender.pdf",
                "location": "page 4",
                "text": "EMD amount is INR 25,000.",
                "detected_values": ["INR 25,000"],
            },
        },
    )

    assert summary["summary"] == "The tender requires an earnest money deposit; amount: INR 25,000."


def test_suggested_action_descriptions_are_grounded_and_fall_back_when_missing():
    evidence = {
        "ELIGIBILITY-01": {
            "category": "eligibility",
            "document": "Tender.pdf",
            "location": "page 4",
            "text": "Bidders must submit proof of five years of relevant experience.",
        }
    }
    result = _validated_document_summaries(
        {
            "analysis": "The eligibility requirement was reviewed.",
            "summaries": [{
                "evidence_id": "ELIGIBILITY-01",
                "summary": "Submit proof of the required five years of relevant experience.",
                "description": "Include evidence covering five years of relevant experience.",
            }],
        },
        evidence,
    )
    action = result["actions"][0]
    assert "five years" in action["description"]
    assert action["evidence_ids"] == ["ELIGIBILITY-01"]

    missing_description = _validated_document_summaries(
        {
            "analysis": "The eligibility requirement was reviewed.",
            "summaries": [{
                "evidence_id": "ELIGIBILITY-01",
                "summary": "Submit proof of the required five years of relevant experience.",
            }],
        },
        evidence,
    )
    assert "does not specify the complete procedure" in missing_description["actions"][0]["description"]

    unsupported_description = _validated_document_summaries(
        {
            "analysis": "The eligibility requirement was reviewed.",
            "summaries": [{
                "evidence_id": "ELIGIBILITY-01",
                "summary": "Submit proof of the required five years of relevant experience.",
                "description": "Submit the form by 31 December.",
            }],
        },
        evidence,
    )
    assert "does not specify the complete procedure" in unsupported_description["actions"][0]["description"]


def test_review_evidence_is_capped_and_keeps_each_category_and_detected_values():
    evidence = {
        f"{category.upper()}-{index:02}": {
            "category": category,
            "text": f"{category} clause {index}",
            **({"detected_values": ["INR 1,000"]} if index == 5 and category == "economics" else {}),
        }
        for category in ("economics", "eligibility", "submissions", "deadlines", "commercial_terms")
        for index in range(1, 6)
    }

    selected = _select_review_evidence(evidence, limit=7)

    assert len(selected) == 7
    assert {item["category"] for item in selected.values()} == {
        "economics", "eligibility", "submissions", "deadlines", "commercial_terms"
    }
    assert "ECONOMICS-05" in selected


def test_model_summary_cannot_add_numbers_from_another_clause():
    assert not _summary_numbers_are_grounded(
        "The tender value is INR 1,25,00,000 for a five-year experience requirement.",
        "Estimated tender value is INR 1,25,00,000.",
    )
    assert _summary_numbers_are_grounded(
        "The tender value is INR 1,25,00,000.",
        "Estimated tender value is INR 1,25,00,000.",
    )


def test_action_plan_leaves_document_based_steps_to_the_live_agent():
    first_clause = {
        "document": "RFP-A.pdf",
        "location": "page 2",
        "excerpt": "Five years of relevant experience is required.",
    }
    second_clause = {
        "document": "RFP-B.pdf",
        "location": "page 4",
        "excerpt": "The bidder must have five years of experience.",
    }
    plan = prepare_pursuit_action_plan(
        tender={"closing_at": "2027-01-30T00:00:00+00:00"},
        qualification={"fit_score": 50, "missing_fields": ["region"]},
        risk_review={"risks": [], "hard_stop": False},
        document_analysis={
            "findings": {
                "eligibility": [first_clause, first_clause, second_clause],
                "submission_requirements": [first_clause, first_clause],
                "deadlines": [second_clause],
            },
        },
    )

    assert plan["recommendation"] == "review"
    assert plan["next_steps"] == []


def test_risk_review_groups_same_risk_and_retains_each_document_reference():
    references = [
        {"document": "Terms.pdf", "location": "page 5", "excerpt": "Late delivery damages apply."},
        {"document": "Schedule.pdf", "location": "page 2", "excerpt": "A penalty applies to delays."},
    ]
    result = review_tender_risks(
        tender={
            "closing_at": "2027-01-30T00:00:00+00:00",
            "existing_win_probability_percent": 75,
        },
        analysis={
            "documents_analyzed": 2,
            "findings": {"commercial_risks": references},
            "explicit_disqualifier_evidence": [],
        },
        evaluated_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )

    commercial_risks = [
        risk for risk in result["risks"]
        if "Commercial clauses" in risk["risk"]
    ]
    assert len(commercial_risks) == 1
    assert commercial_risks[0]["references"] == references


def test_document_evidence_is_cited_and_unsupported_files_fail_explicitly():
    page_data = extract_document_pages(
        "requirements.txt",
        b"Eligibility requires past experience.\n"
        b"A bid deadline applies.\n"
        b"Estimated tender value is INR 1,250,000 and EMD is INR 25,000.\n"
        b"Must submit a price schedule and GST registration certificate.\n"
        b"Submission deadline: 30-12-2026.",
    )
    findings = find_document_evidence(page_data)
    assert findings["findings"]["eligibility"][0]["location"] == "line 1"
    assert findings["findings"]["deadlines"][0]["location"] == "line 2"
    assert findings["findings"]["tender_economics"][0]["location"] == "line 3"
    submission_locations = {
        item["location"] for item in findings["findings"]["submission_requirements"]
    }
    assert {"line 3", "line 4"}.issubset(submission_locations)
    assert [item["value"] for item in findings["extracted_financial_amounts"]] == [
        "INR 1,250,000",
        "INR 25,000",
    ]
    assert findings["extracted_financial_amounts"][0]["document"] == "requirements.txt"
    assert findings["extracted_deadlines"] == [{
        "value": "30-12-2026",
        "document": "requirements.txt",
        "location": "line 5",
        "excerpt": "Submission deadline: 30-12-2026.",
    }]
    with pytest.raises(ValueError, match="supported formats"):
        extract_document_pages("requirements.doc", b"not supported")
    docx_buffer = io.BytesIO()
    docx = Document()
    docx.add_paragraph("Mandatory eligibility certificate.")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Estimated tender value"
    table.cell(0, 1).text = "INR 2,10,00,00,000"
    table.cell(1, 0).text = "Submission deadline"
    table.cell(1, 1).text = "12/12/2026"
    docx.save(docx_buffer)
    docx_pages = extract_document_pages("requirements.docx", docx_buffer.getvalue())
    assert docx_pages[0]["location"] == "paragraph 1"
    assert docx_pages[1] == {
        "filename": "requirements.docx",
        "location": "table 1, row 1",
        "text": "Estimated tender value | INR 2,10,00,00,000",
    }
    docx_findings = find_document_evidence(docx_pages)
    assert docx_findings["extracted_financial_amounts"][0]["location"] == "table 1, row 1"
    assert docx_findings["extracted_financial_amounts"][0]["value"] == "INR 2,10,00,00,000"
    assert docx_findings["extracted_tender_values"] == [{
        "value": "INR 2,10,00,00,000",
        "document": "requirements.docx",
        "location": "table 1, row 1",
        "excerpt": "Estimated tender value | INR 2,10,00,00,000",
    }]
    assert docx_findings["extracted_deadlines"][0]["location"] == "table 1, row 2"


def test_tender_value_and_deadline_extraction_supports_labels_and_times():
    pages = extract_document_pages(
        "tender.txt",
        (
            "Estimated Tender Cost (INR): 2,10,00,00,000\n"
            "Bid submission end date and time: 12/12/2026, 8:47:36 am"
        ).encode(),
    )
    findings = find_document_evidence(pages)
    assert findings["extracted_tender_values"][0]["value"] == "INR 2,10,00,00,000"
    assert findings["extracted_tender_values"][0]["location"] == "line 1"
    assert findings["extracted_deadlines"][0]["value"] == "12/12/2026, 8:47:36 am"
    assert findings["extracted_deadlines"][0]["location"] == "line 2"


def test_document_evidence_classifies_bid_security_as_a_key_tender_cost():
    findings = find_document_evidence([{
        "filename": "security.txt",
        "location": "line 1",
        "text": "Bid security: INR 75,000.",
    }])

    assert findings["extracted_financial_amounts"][0]["cost_type"] == "Bid security"
    assert findings["extracted_financial_amounts"][0]["value"] == "INR 75,000"


def test_document_evidence_extracts_full_economics_timeline_and_numbered_criteria():
    pages = [
        {
            "filename": "works-tender.pdf",
            "location": "page 6",
            "text": (
                "Estimated cost including duties but excluding GST\n"
                "Rs. 10,19,68,127/-\n"
                "Completion Period: 5 Months\n"
                "Earnest Money Deposit (EMD): Rs. 20,19,681/-\n"
                "Tender Fees: Rs. 11,800/-\n"
                "2. TIME SCHEDULE OF TENDER ACTIVITIES\n"
                "1. Invitation of E-NIT 22.12.2023\n"
                "2. Last date for pre-bid queries 28.12.2023 upto 03:00 PM\n"
                "3. Document download deadline 05.01.2024\n"
                "4. Online tender submission deadline 10.01.2024 upto 03:00 PM\n"
                "5. Technical bid opening 11.01.2024 at 03:30 PM\n"
                "5. TECHNICAL PRE-QUALIFICATION CRITERIA\n"
                "5.1 The bidder must have completed similar work worth 80% "
                "of the estimated cost.\n"
                "Alternatively, two similar works worth 50% each are acceptable."
            ),
            "layout_text": (
                "Sl. No  Name of Work  Estimated cost  Completion Period  "
                "Earnest Money Deposit (EMD)  Tender Fees\n"
                "1.  Civil work  Rs. 10,19,68,127/-  5 Months  "
                "Rs. 20,19,681/-  Rs. 11,800/-"
            ),
            "layout_rows": [
                {
                    "y": 340,
                    "cells": [
                        {"x": 200, "text": "Estimated cost"},
                        {"x": 370, "text": "Completion Period"},
                        {"x": 475, "text": "Earnest Money Deposit (EMD)"},
                        {"x": 610, "text": "Tender Fees"},
                    ],
                },
                {
                    "y": 270,
                    "cells": [
                        {"x": 200, "text": "Rs. 10,19,68,127/-"},
                        {"x": 370, "text": "5 Months"},
                        {"x": 475, "text": "Rs. 20,19,681/-"},
                        {"x": 610, "text": "Rs. 11,800/-"},
                    ],
                },
            ],
        },
        {
            "filename": "works-tender.pdf",
            "location": "page 7",
            "text": (
                "Three similar works worth 40% each are also acceptable.\n"
                "5.2 The bidder must meet the billing criterion of 90%.\n"
                "6. SCOPE OF WORK\n"
                "Work details follow."
            ),
        },
    ]

    findings = find_document_evidence(pages)

    assert [item["value"] for item in findings["extracted_financial_amounts"]] == [
        "Rs. 10,19,68,127",
        "Rs. 20,19,681",
        "Rs. 11,800",
    ]
    assert [item.get("cost_type") for item in findings["extracted_financial_amounts"]] == [
        "Estimated cost",
        "Earnest money deposit (EMD)",
        "Tender fee",
    ]
    assert [item["value"] for item in findings["extracted_tender_periods"]] == ["5 Months"]
    assert [item["value"] for item in findings["extracted_timeline"]] == [
        "22 Dec 2023",
        "28 Dec 2023 · 03:00 PM",
        "05 Jan 2024",
        "10 Jan 2024 · 03:00 PM",
        "11 Jan 2024 · 03:30 PM",
    ]
    assert [item["event"] for item in findings["extracted_timeline"]] == [
        "Invitation of E-NIT",
        "Last date for pre-bid queries",
        "Document download deadline",
        "Online tender submission deadline",
        "Technical bid opening",
    ]
    criteria = findings["prequalification_criteria"]
    assert [item["criterion_id"] for item in criteria] == ["5.1", "5.2"]
    assert "80%" in criteria[0]["excerpt"]
    assert "50%" in criteria[0]["excerpt"]
    assert "40%" in criteria[0]["excerpt"]
    assert [reference["location"] for reference in criteria[0]["references"]] == [
        "page 6",
        "page 7",
    ]
    assert "90%" in criteria[1]["excerpt"]


def test_timeline_keeps_only_unique_important_events_and_sorts_by_date():
    schedule = (
        "TIME SCHEDULE OF TENDER ACTIVITIES\n"
        "1. Invitation of E-NIT 22.12.2023\n"
        "2. Tender activity 28.12.2023\n"
        "3. Last date for pre-bid queries 28.12.2023 upto 03:00 PM\n"
        "4. Last Date & Time for downloading of tender documents 04.01.2024 upto 03:00 PM\n"
        "5. Last Date & Time of submission of Tenders (Online) 04.01.2024 upto 03:00 PM\n"
        "6. Date & Time of online opening of Technical Bid 05.01.2024 at 03:00 PM"
    )
    findings = find_document_evidence([
        {
            "filename": "schedule.pdf",
            "location": "page 5",
            "text": schedule,
            "layout_text": schedule,
        },
    ])

    timeline = findings["extracted_timeline"]

    assert [item["event_type"] for item in timeline] == [
        "invitation",
        "pre_bid",
        "document_download",
        "bid_submission",
        "technical_bid_opening",
    ]
    assert [item["value"] for item in timeline] == [
        "22 Dec 2023",
        "28 Dec 2023 · 03:00 PM",
        "04 Jan 2024 · 03:00 PM",
        "04 Jan 2024 · 03:00 PM",
        "05 Jan 2024 · 03:00 PM",
    ]
    assert len({
        (item["event_type"], item["value"])
        for item in timeline
    }) == len(timeline)
    assert "Tender activity" not in " ".join(item["event"] for item in timeline)


def test_timeline_removes_source_list_numbers_from_event_names():
    schedule = (
        "TIME SCHEDULE OF TENDER ACTIVITIES\n"
        "16 Bid submission start date 27.10.2026\n"
        "14 Pre-bid meeting at 02.11.2026\n"
        "15 Post Pre-bid End date Up to 06.11.2026\n"
        "17 Bid submission end date Up to 16.11.2026"
    )
    findings = find_document_evidence([
        {
            "filename": "schedule.pdf",
            "location": "page 13",
            "text": schedule,
        },
    ])

    assert [item["event"] for item in findings["extracted_timeline"]] == [
        "Bid submission start date",
        "Pre-bid meeting at",
        "Post Pre-bid End date Up to",
        "Bid submission end date Up to",
    ]


def test_document_evidence_does_not_truncate_financial_amounts_or_deadlines():
    pages = extract_document_pages(
        "many-costs.txt",
        "\n".join(f"Other cost: Rs. {amount}" for amount in range(1, 36)).encode(),
    )

    findings = find_document_evidence(pages)

    assert len(findings["extracted_financial_amounts"]) == 35


def test_prequalification_summary_uses_exact_source_when_model_summary_is_unsafe():
    source = (
        "Complete similar works worth 80% of estimated cost. Alternatively, "
        "two similar works worth 50% each or three works worth 40% each."
    )

    criteria = [{
        "criterion_id": "5.1",
        "excerpt": source,
        "references": [{"document": "tender.pdf", "location": "page 4", "excerpt": source}],
    }]
    requested, _ = _prequalification_summary_inputs(criteria)
    result = {
        "prequalification_summaries": [{
            "criterion_key": "criterion-1",
            "summary": "Complete works worth 90% of estimated cost.",
        }],
    }
    criteria = _apply_prequalification_summaries(criteria, requested, result)

    assert criteria[0]["summary"] == source
    assert criteria[0]["summary_source"] == "document_excerpt"


def test_prequalification_summary_uses_source_when_summary_is_missing():
    source = "The bidder must meet a turnover threshold of INR 2,00,00,000."

    criteria = [{
        "criterion_id": "5.2",
        "excerpt": source,
        "references": [{"document": "tender.pdf", "location": "page 5", "excerpt": source}],
    }]
    requested, _ = _prequalification_summary_inputs(criteria)
    criteria = _apply_prequalification_summaries(criteria, requested, {})

    assert criteria[0]["summary"] == source
    assert criteria[0]["summary_source"] == "document_excerpt"


def test_prequalification_summary_input_uses_one_bounded_request():
    criteria = [
        {"criterion_id": str(index), "excerpt": "x" * 1000}
        for index in range(1, 21)
    ]

    requested, requested_keys = _prequalification_summary_inputs(criteria)

    assert len(requested) == 10
    assert len(requested_keys) == 10
    assert len(requested) < len(criteria)
    assert len(requested) == len(requested_keys)


@pytest.mark.asyncio
async def test_document_agent_returns_cited_findings_and_next_steps_when_model_times_out(
    monkeypatch: pytest.MonkeyPatch,
):
    async def model_times_out(*_args, **_kwargs):
        raise RuntimeError("document_agent configured-model request exceeded 30 seconds")

    monkeypatch.setattr(
        "agents.pursuit.document_agent.reason_with_live_model",
        model_times_out,
    )
    source = (
        "Estimated tender cost: INR 1,250,000.\n"
        "Eligibility requires five years of relevant experience.\n"
        "Pre-qualification criteria:\n"
        "5.1 The bidder must complete similar work worth 80% of estimated cost."
    )
    state = {
        "tender": {
            "event_id": "model-timeout-test",
            "tender_id": "TEST-1",
            "stage": "Open",
            "material_categories": ["Construction"],
            "existing_win_probability_percent": 70,
            "region": "West",
            "tender_value": 1_250_000,
            "closing_at": "2027-01-30T00:00:00+00:00",
        },
        "active_rules": [],
        "uploaded_documents": [{
            "filename": "real-extracted-source.txt",
            "content": source.encode(),
        }],
        "document_pages": [],
        "agent_trace": [],
        "tool_trace": [],
        "qualification": {
            "fit_score": 70,
            "missing_fields": [],
            "findings": [],
        },
        "risk_review": {
            "risks": [],
            "hard_stop": False,
            "limitations": "Deterministic test setup.",
        },
    }

    await document_agent(state)
    await action_planner_agent(state)

    report = state["document_analysis"]
    assert report["extracted_tender_values"][0]["value"] == "INR 1,250,000"
    assert report["extracted_tender_values"][0]["document"] == "real-extracted-source.txt"
    assert report["prequalification_criteria"][0]["excerpt"] == (
        "The bidder must complete similar work worth 80% of estimated cost."
    )
    assert report["prequalification_criteria"][0]["summary_source"] == "document_excerpt"
    assert report["economics_summaries"]
    assert report["processing_guidelines"]
    assert report["suggested_actions"]
    assert any("timed out" in item for item in report["limitations"])
    assert state["action_plan"]["next_steps"]


def test_rupee_amounts_keep_their_full_cost_clause_and_source_context():
    preference = (
        "Under parallel contracts, the Class-I local Supplier will be preferred "
        "over Class-II local Supplier and Non-local Supplier as per the terms "
        "mentioned in the document; amount: Rs. 10."
    )
    deposit = "Earnest money deposit (EMD): Rs. 12,00,000."
    page_data = extract_document_pages(
        "tender.txt",
        f"{preference}\n{deposit}".encode(),
    )

    findings = find_document_evidence(page_data)

    assert findings["findings"]["tender_economics"][0]["excerpt"] == preference
    amounts = findings["extracted_financial_amounts"]
    assert [(item["value"], item["excerpt"]) for item in amounts] == [
        ("Rs. 10", preference),
        ("Rs. 12,00,000", deposit),
    ]

    split_layout_findings = find_document_evidence([{
        "filename": "layout.pdf",
        "location": "page 1",
        "text": (
            "Estimated Tender Value\n"
            "(INR)\n"
            "2,10,00,00,000\n"
            "Bid submission end date and time\n"
            "12/12/2026, 8:47:36 am"
        ),
    }])
    assert split_layout_findings["extracted_tender_values"][0]["value"] == "INR 2,10,00,00,000"
    assert split_layout_findings["extracted_deadlines"][0]["value"] == "12/12/2026, 8:47:36 am"


def test_tender_value_label_does_not_claim_nearby_preference_clause_amount():
    findings = find_document_evidence([{
        "filename": "TD.pdf",
        "location": "page 9",
        "text": (
            "Tender value\n"
            "Under parallel contracts, the Class-I local Supplier will be preferred "
            "over Class-II local Supplier as per document terms; amount: Rs. 10."
        ),
    }])

    assert findings["extracted_tender_values"] == []
    assert findings["extracted_financial_amounts"] == [{
        "value": "Rs. 10",
        "document": "TD.pdf",
        "location": "page 9",
        "excerpt": (
            "Under parallel contracts, the Class-I local Supplier will be preferred "
            "over Class-II local Supplier as per document terms; amount: Rs. 10."
        ),
    }]


def test_non_exclusion_eligibility_wording_is_not_flagged_as_critical():
    from agents.tools.pursuit_review_tools import review_tender_risks

    pages = [{
        "filename": "eligibility.pdf",
        "location": "page 10",
        "text": (
            "The Bidder should not have been blacklisted or Debarred in any State Govt./Municipal. "
            "The tenderers are liable to be disqualified and not considered for the tender."
        ),
    }]
    analysis = find_document_evidence(pages)
    exclusion_excerpts = analysis["explicit_disqualifier_evidence"]
    assert len(exclusion_excerpts) == 1
    assert "liable to be disqualified" in exclusion_excerpts[0]["excerpt"]
    assert "should not have been blacklisted" not in exclusion_excerpts[0]["excerpt"]
    risk_review = review_tender_risks(
        {
            "closing_at": (datetime.now(timezone.utc) + timedelta(days=30)).isoformat(),
            "existing_win_probability_percent": 70,
        },
        analysis,
    )
    document_risks = [
        risk for risk in risk_review["risks"]
        if any(reference["document"] == "eligibility.pdf" for reference in risk.get("references", []))
    ]
    assert len(document_risks) == 1
    assert document_risks[0]["level"] == "critical"


@pytest.mark.asyncio
async def test_copilot_rejects_scanned_pdf_and_non_tenant_tender(copilot_client):
    client, session_factory, _ = copilot_client
    async with session_factory() as session:
        session.add(TenderEvent(
            tenant_id="TENANT-B",
            event_id="another-tenant-tender",
            sender_entity_id="TENANT-B",
            tender_ref="RFP-OTHER",
            project_name="Private tender",
            source_revision="rev-1",
            prior_stage=None,
            stage="Reviewing",
            occurred_at=datetime.now(timezone.utc),
            material_categories=["Steel"],
            expected_need_window=None,
            existing_win_probability_percent=None,
            probability_updated_at=None,
            source_type="crm",
            data_provenance="test",
        ))
        await session.commit()

    async with session_factory() as session:
        session.add(TenderEvent(
            tenant_id="TENANT-A",
            event_id="own-blank-pdf",
            sender_entity_id="TENANT-A",
            tender_ref="RFP-BLANK",
            project_name="Blank PDF",
            source_revision="rev-1",
            prior_stage=None,
            stage="Reviewing",
            occurred_at=datetime.now(timezone.utc),
            material_categories=["Steel"],
            expected_need_window=None,
            existing_win_probability_percent=None,
            probability_updated_at=None,
            source_type="manual",
            data_provenance="test",
        ))
        await session.commit()

    missing = await client.post(
        "/api/v1/pursuit/another-tenant-tender/analyze",
        files=[("files", ("brief.pdf", b"%PDF-1.4 no extractable content", "application/pdf"))],
    )
    assert missing.status_code == 404

    malformed = await client.post(
        "/api/v1/pursuit/own-blank-pdf/analyze",
        files=[("files", ("broken.pdf", b"not a PDF", "application/pdf"))],
    )
    assert malformed.status_code == 422
    assert "unable to read this PDF" in malformed.json()["detail"]

    pdf_buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(pdf_buffer)
    own_tender = await client.post(
        "/api/v1/pursuit/own-blank-pdf/analyze",
        files=[("files", ("brief.pdf", pdf_buffer.getvalue(), "application/pdf"))],
    )
    assert own_tender.status_code == 422
    assert "scanned PDFs need OCR" in own_tender.json()["detail"]
