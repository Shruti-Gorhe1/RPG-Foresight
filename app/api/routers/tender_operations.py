"""Tender import, management, feedback, follow-up, and reporting APIs."""
from __future__ import annotations

import csv
import io
import json
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any, Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from api.auth import CurrentUser, get_current_user, require_tenant
from db.base import get_db_session
from db.models import (
    RoutingRule,
    RecipientAddress,
    NotificationDelivery,
    SignalFeedback,
    TenderEvent,
    TenderFollowUp,
    TenderOutcomeSnapshot,
)
from services.win_probability import (
    apply_win_probability,
    refresh_tender_win_probability,
)

router = APIRouter()
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 2_000
FOLLOW_UP_STATUSES = {"New", "Reviewing", "Bid / No bid", "Closed"}
FEEDBACK_VALUES = {"Relevant", "Not relevant", "Already known"}


class TenderInput(BaseModel):
    event_id: str | None = Field(default=None, max_length=64)
    tender_id: str = Field(min_length=1, max_length=64)
    project_name: str = Field(min_length=1, max_length=128)
    stage: str = Field(min_length=1, max_length=32)
    material_categories: list[str] = Field(min_length=1)
    existing_win_probability_percent: float | None = Field(
        default=None,
        ge=0,
        le=100,
        description="Deprecated legacy input; win probability is calculated from historical outcomes.",
    )
    region: str | None = Field(default=None, max_length=128)
    tender_value: float | None = Field(default=None, ge=0)
    currency: str | None = Field(default=None, min_length=3, max_length=8)
    opening_at: datetime | None = None
    closing_at: datetime | None = None


class FeedbackInput(BaseModel):
    feedback: Literal["Relevant", "Not relevant", "Already known"]
    reason: str | None = Field(default=None, max_length=1000)


class FollowUpInput(BaseModel):
    assigned_to: str | None = Field(default=None, max_length=254)
    status: Literal["New", "Reviewing", "Bid / No bid", "Closed"] = "New"
    notes: str = Field(default="", max_length=10_000)
    follow_up_at: datetime | None = None


def _serialize_tender(event: TenderEvent) -> dict[str, Any]:
    return {
        "id": str(event.id),
        "event_id": event.event_id,
        "tender_id": event.tender_ref,
        "project_name": event.project_name,
        "stage": event.stage,
        "prior_stage": event.prior_stage,
        "source_type": event.source_type,
        "data_provenance": event.data_provenance,
        "existing_win_probability_percent": event.existing_win_probability_percent,
        "probability_updated_at": (
            event.probability_updated_at.isoformat()
            if event.probability_updated_at
            else None
        ),
        "material_categories": event.material_categories or [],
        "region": event.region,
        "tender_value": event.tender_value,
        "currency": event.currency,
        "opening_at": event.opening_at.isoformat() if event.opening_at else None,
        "closing_at": event.closing_at.isoformat() if event.closing_at else None,
        "source_updated_at": event.occurred_at.isoformat() if event.occurred_at else None,
        "ingested_at": event.ingested_at.isoformat() if event.ingested_at else None,
        "sender_entity_id": event.sender_entity_id,
        "tenant_id": event.tenant_id,
        "source_revision": event.source_revision,
    }


def _parse_datetime(value: Any) -> datetime | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def _as_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _read_upload(filename: str | None, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Upload is larger than the 10 MB limit")
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="CSV must be UTF-8 encoded") from exc
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(status_code=422, detail="CSV is missing a header row")
        headers = [header.strip() for header in reader.fieldnames if header]
        if len(headers) != len(set(headers)):
            raise HTTPException(status_code=422, detail="CSV headers must be unique")
        rows = [{str(key).strip(): value for key, value in row.items() if key} for row in reader]
    elif suffix in {"xlsx", "xlsm"}:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            raw_headers = next(values, None)
            if not raw_headers:
                raise HTTPException(status_code=422, detail="Spreadsheet is missing a header row")
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            nonempty_headers = [header for header in headers if header]
            if len(nonempty_headers) != len(set(nonempty_headers)):
                raise HTTPException(status_code=422, detail="Spreadsheet headers must be unique")
            rows = [
                {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
                for row in values
                if any(value is not None for value in row)
            ]
            workbook.close()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Unable to read the Excel workbook") from exc
    else:
        raise HTTPException(status_code=415, detail="Upload a .csv, .xlsx, or .xlsm file")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=413, detail=f"Import is limited to {MAX_IMPORT_ROWS} rows")
    return headers, rows


def _mapped_tender(row: dict[str, Any], mapping: dict[str, str]) -> dict[str, Any]:
    def get(field: str) -> Any:
        column = mapping.get(field)
        return row.get(column) if column else None

    materials_value = get("material_categories")
    if isinstance(materials_value, (list, tuple)):
        materials = [str(value).strip() for value in materials_value if str(value).strip()]
    else:
        materials = [item.strip() for item in str(materials_value or "").replace("|", ",").replace(";", ",").split(",") if item.strip()]
    probability = get("existing_win_probability_percent")
    tender_value = get("tender_value")
    return {
        "event_id": str(get("event_id") or "").strip() or None,
        "tender_id": str(get("tender_id") or "").strip(),
        "project_name": str(get("project_name") or "").strip(),
        "stage": str(get("stage") or "").strip(),
        "material_categories": materials,
        "existing_win_probability_percent": float(probability) if probability is not None and str(probability).strip() else None,
        "region": str(get("region") or "").strip() or None,
        "tender_value": float(tender_value) if tender_value is not None and str(tender_value).strip() else None,
        "currency": str(get("currency") or "").strip().upper() or None,
        "opening_at": _parse_datetime(get("opening_at")),
        "closing_at": _parse_datetime(get("closing_at")),
    }


async def _filtered_events(
    db: AsyncSession,
    tenant_id: str,
    *,
    search: str | None = None,
    stage: str | None = None,
    source_type: str | None = None,
    region: str | None = None,
    material: str | None = None,
    currency: str | None = None,
    probability_min: float | None = None,
    closing_before: datetime | None = None,
    minimum_value: float | None = None,
    maximum_value: float | None = None,
) -> list[TenderEvent]:
    query = select(TenderEvent).where(TenderEvent.tenant_id == tenant_id)
    if stage:
        query = query.where(TenderEvent.stage == stage)
    if source_type:
        query = query.where(TenderEvent.source_type == source_type)
    if region:
        query = query.where(TenderEvent.region.ilike(f"%{region}%"))
    if currency:
        query = query.where(TenderEvent.currency.ilike(currency.strip()))
    if probability_min is not None:
        query = query.where(TenderEvent.existing_win_probability_percent >= probability_min)
    if closing_before:
        query = query.where(TenderEvent.closing_at <= closing_before)
    if minimum_value is not None:
        query = query.where(TenderEvent.tender_value >= minimum_value)
    if maximum_value is not None:
        query = query.where(TenderEvent.tender_value <= maximum_value)
    result = await db.execute(query.order_by(TenderEvent.ingested_at.desc()).limit(5000))
    events = result.scalars().all()
    if search:
        text = search.casefold()
        events = [e for e in events if text in e.project_name.casefold() or text in e.tender_ref.casefold()]
    if material:
        material_key = material.casefold()
        events = [e for e in events if any(material_key in str(item).casefold() for item in (e.material_categories or []))]
    return events


@router.get("")
async def list_tenders(
    search: str | None = None,
    stage: str | None = None,
    source_type: str | None = None,
    region: str | None = None,
    material: str | None = None,
    currency: str | None = None,
    probability_min: float | None = Query(default=None, ge=0, le=100),
    closing_before: datetime | None = None,
    minimum_value: float | None = Query(default=None, ge=0),
    maximum_value: float | None = Query(default=None, ge=0),
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> list[dict[str, Any]]:
    tenant_id = require_tenant(current_user)
    events = await _filtered_events(
        db, tenant_id, search=search, stage=stage, source_type=source_type,
        region=region, material=material, currency=currency, probability_min=probability_min,
        closing_before=closing_before, minimum_value=minimum_value, maximum_value=maximum_value,
    )
    return [_serialize_tender(event) for event in events]


@router.post("", status_code=status.HTTP_201_CREATED)
async def create_tender(
    payload: TenderInput,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    event_id = payload.event_id or f"manual-{uuid.uuid4().hex}"
    duplicate = await db.execute(select(TenderEvent.id).where(TenderEvent.event_id == event_id))
    if duplicate.scalar_one_or_none() is not None:
        raise HTTPException(status_code=409, detail="Tender event ID already exists")
    now = datetime.now(timezone.utc)
    event = TenderEvent(
        event_id=event_id,
        tenant_id=tenant_id,
        sender_entity_id=tenant_id,
        tender_ref=payload.tender_id,
        project_name=payload.project_name,
        source_revision=f"manual-{uuid.uuid4().hex[:12]}",
        prior_stage=None,
        stage=payload.stage,
        occurred_at=now,
        material_categories=payload.material_categories,
        expected_need_window=None,
        existing_win_probability_percent=None,
        probability_updated_at=None,
        source_type="manual",
        data_provenance=f"manual entry by {current_user.user_id}",
        region=payload.region,
        tender_value=payload.tender_value,
        currency=payload.currency,
        opening_at=_as_utc(payload.opening_at),
        closing_at=_as_utc(payload.closing_at),
    )
    db.add(event)
    await db.flush()
    await refresh_tender_win_probability(db, event)
    await db.commit()
    await db.refresh(event)
    return _serialize_tender(event)


@router.patch("/{event_id}")
async def update_tender(
    event_id: str,
    payload: TenderInput,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    result = await db.execute(
        select(TenderEvent).where(TenderEvent.event_id == event_id, TenderEvent.tenant_id == tenant_id)
    )
    event = result.scalar_one_or_none()
    if event is None:
        raise HTTPException(status_code=404, detail="Tender not found")
    updates = payload.model_dump(exclude_unset=True)
    field_map = {
        "tender_id": "tender_ref",
    }
    for field, value in updates.items():
        if field in {"event_id", "existing_win_probability_percent"}:
            continue
        if field in {"opening_at", "closing_at"}:
            value = _as_utc(value)
        setattr(event, field_map.get(field, field), value)
    event.source_revision = f"edited-{uuid.uuid4().hex[:12]}"
    await refresh_tender_win_probability(db, event)
    await db.commit()
    await db.refresh(event)
    return _serialize_tender(event)


@router.post("/import/preview")
async def preview_import(file: UploadFile = File(...)) -> dict[str, Any]:
    headers, rows = _read_upload(file.filename, await file.read(MAX_UPLOAD_BYTES + 1))
    return {"headers": headers, "preview": rows[:10], "row_count": len(rows), "maximum_rows": MAX_IMPORT_ROWS}


@router.post("/import")
async def import_tenders(
    file: UploadFile = File(...),
    mapping_json: str = Form(...),
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    try:
        mapping = json.loads(mapping_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=422, detail="Field mapping must be valid JSON") from exc
    allowed_fields = {
        "event_id", "tender_id", "project_name", "stage", "material_categories",
        "region", "tender_value", "currency", "opening_at", "closing_at",
    }
    if not isinstance(mapping, dict) or any(key not in allowed_fields for key in mapping):
        raise HTTPException(status_code=422, detail="Field mapping contains an unsupported field")
    required_fields = {"tender_id", "project_name", "stage", "material_categories"}
    if not required_fields.issubset(mapping) or any(not isinstance(value, str) or not value for value in mapping.values()):
        raise HTTPException(status_code=422, detail="Map tender ID, project name, stage, and materials to source columns")
    headers, rows = _read_upload(file.filename, await file.read(MAX_UPLOAD_BYTES + 1))
    if any(column not in headers for column in mapping.values()):
        raise HTTPException(status_code=422, detail="Field mapping references a column that is not in the uploaded file")
    imported = 0
    rejected: list[dict[str, Any]] = []
    seen_event_ids: set[str] = set()
    history_result = await db.execute(
        select(TenderOutcomeSnapshot).where(
            TenderOutcomeSnapshot.tenant_id == tenant_id,
            TenderOutcomeSnapshot.is_demo.is_(False),
        )
    )
    outcome_snapshots = list(history_result.scalars().all())
    for row_number, row in enumerate(rows, start=2):
        try:
            values = _mapped_tender(row, mapping)
            data = TenderInput(**values)
            if not data.material_categories:
                raise ValueError("At least one material category is required")
            event_id = data.event_id or f"import-{uuid.uuid4().hex}"
            if event_id in seen_event_ids:
                raise ValueError(f"Event ID '{event_id}' appears more than once in this upload")
            existing = await db.execute(select(TenderEvent.id).where(TenderEvent.event_id == event_id))
            if existing.scalar_one_or_none() is not None:
                raise ValueError(f"Event ID '{event_id}' already exists")
            seen_event_ids.add(event_id)
            now = datetime.now(timezone.utc)
            event = TenderEvent(
                event_id=event_id, tenant_id=tenant_id, sender_entity_id=tenant_id,
                tender_ref=data.tender_id, project_name=data.project_name,
                source_revision=f"import-{uuid.uuid4().hex[:12]}", prior_stage=None,
                stage=data.stage, occurred_at=now, material_categories=data.material_categories,
                expected_need_window=None,
                existing_win_probability_percent=None,
                probability_updated_at=None,
                source_type="file_import", data_provenance=f"file import by {current_user.user_id}",
                region=data.region, tender_value=data.tender_value, currency=data.currency,
                opening_at=data.opening_at, closing_at=data.closing_at,
            )
            apply_win_probability(event, outcome_snapshots)
            db.add(event)
            imported += 1
        except (ValueError, TypeError, OverflowError) as exc:
            rejected.append({"row": row_number, "reason": str(exc)})
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Import could not be saved because one or more tender IDs already exist") from exc
    return {"imported": imported, "rejected_count": len(rejected), "rejected_rows": rejected}


@router.put("/{event_id}/feedback")
async def save_feedback(
    event_id: str,
    payload: FeedbackInput,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    event_result = await db.execute(
        select(TenderEvent.id).where(TenderEvent.event_id == event_id, TenderEvent.tenant_id == tenant_id)
    )
    if event_result.scalar_one_or_none() is None:
        raise HTTPException(status_code=404, detail="Tender not found")
    result = await db.execute(
        select(SignalFeedback).where(
            SignalFeedback.event_id == event_id,
            SignalFeedback.tenant_id == tenant_id,
            SignalFeedback.user_id == current_user.user_id,
        )
    )
    feedback = result.scalar_one_or_none()
    if feedback is None:
        feedback = SignalFeedback(event_id=event_id, tenant_id=tenant_id, user_id=current_user.user_id, feedback=payload.feedback)
        db.add(feedback)
    feedback.feedback = payload.feedback
    feedback.reason = payload.reason
    feedback.created_at = datetime.now(timezone.utc)
    await db.commit()
    return {"event_id": event_id, "feedback": feedback.feedback, "reason": feedback.reason}


@router.get("/{event_id}/feedback")
async def get_feedback(
    event_id: str,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    result = await db.execute(
        select(SignalFeedback).where(
            SignalFeedback.event_id == event_id,
            SignalFeedback.tenant_id == tenant_id,
            SignalFeedback.user_id == current_user.user_id,
        )
    )
    feedback = result.scalar_one_or_none()
    return {
        "event_id": event_id,
        "feedback": feedback.feedback if feedback else None,
        "reason": feedback.reason if feedback else None,
    }


@router.get("/feedback/suggestions")
async def feedback_suggestions(
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    result = await db.execute(select(SignalFeedback).where(SignalFeedback.tenant_id == tenant_id))
    entries = result.scalars().all()
    counts = {value: 0 for value in FEEDBACK_VALUES}
    for entry in entries:
        counts[entry.feedback] = counts.get(entry.feedback, 0) + 1
    suggestions = []
    if counts["Not relevant"]:
        suggestions.append("Review stage, region, material, and probability filters for tenders marked Not relevant.")
    if counts["Already known"]:
        suggestions.append("Consider tightening freshness limits or checking source revisions for tenders marked Already known.")
    if counts["Relevant"] > 0 and counts["Not relevant"] == 0:
        suggestions.append("Current feedback is positive; keep the enabled routing criteria under review.")
    return {"counts": counts, "recommendations": suggestions, "automatic_rule_changes": False}


@router.get("/follow-ups")
async def list_follow_ups(
    due_only: bool = False,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> list[dict[str, Any]]:
    tenant_id = require_tenant(current_user)
    result = await db.execute(
        select(TenderFollowUp, TenderEvent)
        .join(TenderEvent, TenderEvent.event_id == TenderFollowUp.event_id)
        .where(TenderFollowUp.tenant_id == tenant_id, TenderEvent.tenant_id == tenant_id)
        .order_by(TenderFollowUp.follow_up_at.asc().nullslast())
    )
    now = datetime.now(timezone.utc)
    output = []
    for follow_up, event in result.all():
        due = follow_up.follow_up_at is not None and (
            follow_up.follow_up_at.replace(tzinfo=timezone.utc)
            if follow_up.follow_up_at.tzinfo is None
            else follow_up.follow_up_at
        ) <= now
        if due_only and not due:
            continue
        output.append({
            "event_id": follow_up.event_id,
            "tender_id": event.tender_ref,
            "project_name": event.project_name,
            "closing_at": event.closing_at.isoformat() if event.closing_at else None,
            "assigned_to": follow_up.assigned_to,
            "status": follow_up.status,
            "notes": follow_up.notes,
            "follow_up_at": follow_up.follow_up_at.isoformat() if follow_up.follow_up_at else None,
            "is_due": due,
            "updated_at": follow_up.updated_at.isoformat() if follow_up.updated_at else None,
        })
    return output


@router.put("/{event_id}/follow-up")
async def save_follow_up(
    event_id: str,
    payload: FollowUpInput,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    event_result = await db.execute(
        select(TenderEvent).where(TenderEvent.event_id == event_id, TenderEvent.tenant_id == tenant_id)
    )
    event = event_result.scalar_one_or_none()
    if event is None:
        raise HTTPException(status_code=404, detail="Tender not found")
    result = await db.execute(select(TenderFollowUp).where(TenderFollowUp.event_id == event_id, TenderFollowUp.tenant_id == tenant_id))
    follow_up = result.scalar_one_or_none()
    if follow_up is None:
        follow_up = TenderFollowUp(event_id=event_id, tenant_id=tenant_id, updated_by=current_user.user_id)
        db.add(follow_up)
    follow_up.assigned_to = payload.assigned_to
    follow_up.status = payload.status
    follow_up.notes = payload.notes
    follow_up.follow_up_at = payload.follow_up_at
    follow_up.updated_by = current_user.user_id
    await db.commit()
    return {"event_id": event_id, **payload.model_dump(mode="json")}


@router.get("/reports/summary")
async def report_summary(
    search: str | None = None,
    stage: str | None = None,
    source_type: str | None = None,
    region: str | None = None,
    material: str | None = None,
    currency: str | None = None,
    probability_min: float | None = Query(default=None, ge=0, le=100),
    closing_before: datetime | None = None,
    minimum_value: float | None = Query(default=None, ge=0),
    maximum_value: float | None = Query(default=None, ge=0),
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    tenant_id = require_tenant(current_user)
    events = await _filtered_events(
        db, tenant_id, search=search, stage=stage, source_type=source_type,
        region=region, material=material, currency=currency, probability_min=probability_min,
        closing_before=closing_before, minimum_value=minimum_value, maximum_value=maximum_value,
    )
    by_stage: dict[str, int] = {}
    by_material: dict[str, int] = {}
    by_source: dict[str, int] = {}
    by_region: dict[str, int] = {}
    by_probability: dict[str, int] = {"0-24": 0, "25-49": 0, "50-74": 0, "75-100": 0, "Unspecified": 0}
    by_recipient: dict[str, int] = {}
    by_deadline: dict[str, int] = {"Overdue": 0, "0-7 days": 0, "8-30 days": 0, "More than 30 days": 0, "No closing date": 0}
    due_soon = 0
    now = datetime.now(timezone.utc)
    for event in events:
        by_stage[event.stage] = by_stage.get(event.stage, 0) + 1
        by_source[event.source_type] = by_source.get(event.source_type, 0) + 1
        if event.region:
            by_region[event.region] = by_region.get(event.region, 0) + 1
        for value in event.material_categories or []:
            by_material[value] = by_material.get(value, 0) + 1
        probability = event.existing_win_probability_percent
        bucket = "Unspecified" if probability is None else "0-24" if probability < 25 else "25-49" if probability < 50 else "50-74" if probability < 75 else "75-100"
        by_probability[bucket] += 1
        closing = event.closing_at
        if closing is not None:
            closing = closing.replace(tzinfo=timezone.utc) if closing.tzinfo is None else closing
            if closing < now:
                by_deadline["Overdue"] += 1
            elif closing <= now + timedelta(days=7):
                by_deadline["0-7 days"] += 1
                due_soon += 1
            elif closing <= now + timedelta(days=30):
                by_deadline["8-30 days"] += 1
                due_soon += 1
            else:
                by_deadline["More than 30 days"] += 1
        else:
            by_deadline["No closing date"] += 1
    deliveries = await db.execute(
        select(NotificationDelivery.recipient_user_id)
        .join(TenderEvent, TenderEvent.event_id == NotificationDelivery.event_id)
        .where(TenderEvent.tenant_id == tenant_id)
    )
    for (recipient,) in deliveries.all():
        by_recipient[recipient] = by_recipient.get(recipient, 0) + 1
    followups = await list_follow_ups(db=db, current_user=current_user)
    return {
        "total_tenders": len(events),
        "closing_within_30_days": due_soon,
        "due_follow_ups": sum(1 for item in followups if item["is_due"]),
        "by_stage": by_stage,
        "by_material": by_material,
        "by_source": by_source,
        "by_region": by_region,
        "by_probability": by_probability,
        "by_deadline": by_deadline,
        "by_recipient": by_recipient,
    }


@router.get("/reports/export.csv")
async def export_tenders(
    search: str | None = None,
    stage: str | None = None,
    source_type: str | None = None,
    region: str | None = None,
    material: str | None = None,
    currency: str | None = None,
    probability_min: float | None = Query(default=None, ge=0, le=100),
    closing_before: datetime | None = None,
    minimum_value: float | None = Query(default=None, ge=0),
    maximum_value: float | None = Query(default=None, ge=0),
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> StreamingResponse:
    tenant_id = require_tenant(current_user)
    events = await _filtered_events(
        db, tenant_id, search=search, stage=stage, source_type=source_type,
        region=region, material=material, currency=currency, probability_min=probability_min,
        closing_before=closing_before, minimum_value=minimum_value, maximum_value=maximum_value,
    )
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=[
        "event_id", "tender_id", "project_name", "stage", "materials", "region",
        "tender_value", "currency", "probability_percent", "opening_at", "closing_at", "source_type",
    ])
    writer.writeheader()
    for event in events:
        writer.writerow({
            "event_id": event.event_id, "tender_id": event.tender_ref,
            "project_name": event.project_name, "stage": event.stage,
            "materials": ", ".join(event.material_categories or []), "region": event.region or "",
            "tender_value": event.tender_value if event.tender_value is not None else "",
            "currency": event.currency or "",
            "probability_percent": event.existing_win_probability_percent if event.existing_win_probability_percent is not None else "",
            "opening_at": event.opening_at.isoformat() if event.opening_at else "",
            "closing_at": event.closing_at.isoformat() if event.closing_at else "",
            "source_type": event.source_type,
        })
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="tender-report.csv"'},
    )


@router.post("/rules/{rule_id}/preview")
async def preview_rule(
    rule_id: str,
    db: AsyncSession = Depends(get_db_session),
    current_user: CurrentUser = Depends(get_current_user),
) -> dict[str, Any]:
    from services.routing_service import _evaluate_event_for_rule, _latest_tender_events_statement

    tenant_id = require_tenant(current_user)
    rule_result = await db.execute(
        select(RoutingRule).options(selectinload(RoutingRule.routes)).where(
            RoutingRule.rule_id == rule_id, RoutingRule.tenant_id == tenant_id
        )
    )
    rule = rule_result.scalar_one_or_none()
    if rule is None:
        raise HTTPException(status_code=404, detail="Rule not found")
    events_result = await db.execute(
        _latest_tender_events_statement(tenant_id).limit(500)
    )
    events = events_result.scalars().all()
    recipient_result = await db.execute(
        select(RecipientAddress.email).where(RecipientAddress.tenant_id == tenant_id)
    )
    recipient_emails = list(recipient_result.scalars().all())
    matches = 0
    reasons: dict[str, int] = {}
    samples = []
    original_enabled = rule.enabled
    rule.enabled = True
    for event in events:
        evaluation = _evaluate_event_for_rule(event, rule, recipient_emails)
        matched = evaluation.status == "candidate" or any(
            decision.disposition == "candidate" for decision in evaluation.recipient_decisions
        )
        if matched:
            matches += 1
        else:
            reason = evaluation.reason
            if evaluation.recipient_decisions:
                reason = evaluation.recipient_decisions[0].reason
            reasons[reason] = reasons.get(reason, 0) + 1
        if len(samples) < 100:
            sample_reason = "matched" if matched else evaluation.reason
            if not matched and evaluation.recipient_decisions:
                sample_reason = evaluation.recipient_decisions[0].reason
            samples.append({
                "event_id": event.event_id, "tender_id": event.tender_ref,
                "project_name": event.project_name, "match": matched,
                "reason": sample_reason,
            })
    rule.enabled = original_enabled
    rule.previewed_at = datetime.now(timezone.utc)
    await db.commit()
    return {"rule_id": rule.rule_id, "evaluated": len(events), "matches": matches, "excluded": len(events) - matches, "exclusion_reasons": reasons, "samples": samples}
